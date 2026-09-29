"""The Estimate tab: the owner's own crop estimate for a season, block by
block, set against each block's history and - once picking starts - against
what is actually coming in.

This is how the farm has always estimated, moved out of the "Produksie"
sheet of its OES workbook: walk the orchard, judge each block's kg per tree
against what that block did in earlier seasons, multiply by its trees, add
up. The app does not make the call. It lays out the history the call is
made against (last season, the last five seasons, the best five of the last
ten), offers "last season +12%" style fills to start from, keeps every
version of the estimate, and during the season shows how the picking is
tracking against it.

It is deliberately separate from the Risk tab's Harvest Forecast
(routers/risk.py). That one is a whole-farm figure driven by weather alone;
this one is the owner's judgement per block. The two are worth reading side
by side precisely because they are built from different things.

The in-season projection is the one number here the app computes itself: it
scales the kg picked so far by the share of a season that had typically
been picked by the same day in the daily-tracked seasons on file. Early in
a season that share is small and swings wildly from year to year (a late
start is exactly the case where it misleads), so it is only offered once
the typical season is at least MIN_PROJECTION_SHARE picked, and always with
the range the individual seasons give - leaving out, by name, any season
that had barely started by the same day (under MIN_RANGE_SHARE), whose
share would put the top of the range in the thousands of tonnes.

Two more things hang off a version:

  * a pack-out mix (YieldEstimatePack) - the owner's shares of the net
    picked kg per channel and pack type, turned into season kg and cartons
    the way the farm's "%" sheet does (kg x share / kg per carton).
    Estimation only: pallets, transport and the markets are not this app's.
  * a snapshot of the Risk tab's weather-driven Harvest Forecast as it was
    when the version was saved, so a season can later be read back as "my
    estimate, the weather model at the time, and what was picked". This
    module never runs that model itself - the browser sends the figures it
    was showing (see EstimateForecastIn) - so saving never waits on the
    weather scan or an outbound fetch.

Similar past seasons live in routers/analogs.py, loaded separately: they
read decades of weather and must not slow this tab's editor down.
"""
import io
import math
import os
from datetime import date, datetime, timedelta, timezone
from typing import List, Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl.styles import Font
from pydantic import BaseModel, Field as PydField
from sqlmodel import Session, delete, select

from db import get_boord_session, get_owner_session, own_farm_block_ids
from models_boord import Block, SystemSetting
from models_owner import (HistoricalAnnualYield, YieldEstimate, YieldEstimateBlock,
                          YieldEstimatePack)
from routers.analysis import _block_sort_key, block_season_kg, season_day_kg
from routers.historical_report import REPORTS_DIR, XLSX_MEDIA, _style_header_cell
from timeutil import season_day

router = APIRouter(prefix="/api/estimate", tags=["estimate"])

# How far back the per-block reference figures look. Ten seasons is the
# window the farm's own "5 goeie uit 10" method uses; the recent average is
# the last five of those.
REFERENCE_SEASONS = 10
RECENT_SEASONS = 5
BEST_OF = 5

# The in-season projection is withheld until the typical season on file was
# at least this far picked by the same day. Below it, dividing by a small
# and very variable share turns a few days' picking into nonsense.
MIN_PROJECTION_SHARE = 0.25
# A season that had picked less than this by the same day is left out of the
# projection's range: it had barely started, and dividing by it gives an
# upper bound in the thousands of tonnes.
MIN_RANGE_SHARE = 0.10

_SEASON_MIN, _SEASON_MAX = 1980, 2100


def _utcnow() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _r(x, nd=1):
    return round(x, nd) if x is not None else None


def _mean(xs):
    return sum(xs) / len(xs) if xs else None


# --------------------------------------------------------------------------- #
# Reference figures
# --------------------------------------------------------------------------- #
def _block_reference(by_year: dict, trees: int, season_year: int, annual_only: set,
                     current_year: int) -> dict:
    """One block's history and the kg/tree figures an estimate is judged
    against. kg/tree uses the block's CURRENT tree count for every season,
    the same as the Analysis tab - Boord keeps no tree-count history."""
    def kg_tree(kg):
        return kg / trees if trees else None

    history = {}
    for year, kg in by_year.items():
        if season_year - REFERENCE_SEASONS <= year < season_year:
            history[year] = {"kg": _r(kg), "kg_tree": _r(kg_tree(kg)),
                             "annual_only": year in annual_only,
                             "partial": year == current_year}

    # Only finished seasons with fruit count toward the averages: a season
    # still being picked would drag them down, and a zero is a block that
    # wasn't bearing, not a bad crop.
    done = {y: kg for y, kg in by_year.items()
            if y < season_year and y != current_year and kg > 0}
    window = [kg_tree(kg) for y, kg in done.items() if y >= season_year - REFERENCE_SEASONS]
    recent = [kg_tree(kg) for y, kg in done.items() if y >= season_year - RECENT_SEASONS]
    window = [v for v in window if v is not None]
    recent = [v for v in recent if v is not None]
    last = done.get(season_year - 1)
    return {
        "history": history,
        "last_kg_tree": _r(kg_tree(last)) if last is not None else None,
        "avg5_kg_tree": _r(_mean(recent)),
        "best5_kg_tree": _r(_mean(sorted(window, reverse=True)[:BEST_OF])),
        "low_kg_tree": _r(min(window)) if window else None,
        "high_kg_tree": _r(max(window)) if window else None,
        "seasons_on_file": len(window),
    }


def _farm_totals(owner: Session, season: dict) -> tuple:
    """Whole-farm kg per season, every season on file, and the set of
    seasons that have a per-block breakdown. Built the way the Risk tab
    builds its season totals (risk._compute_driver_state): the daily record
    where a season has one, else that season's annual rows - per block
    where it has them, whichever blocks they name (a block pulled out since
    still grew that season's fruit), the single whole-farm figure before
    that (1987-2009). block_season_kg() is deliberately NOT the source here:
    it keeps only blocks in today's register, which is right for per-block
    reference figures and wrong for what the farm picked. Shared with
    routers/analogs.py."""
    totals: dict = {}
    for (y, _, _), kg in season["day_kg"].items():
        totals[y] = totals.get(y, 0.0) + kg
    per_block_years = set(totals)
    annual_block: dict = {}
    annual_farm: dict = {}
    for a in owner.exec(select(HistoricalAnnualYield)).all():
        target = annual_block if a.block_id else annual_farm
        target[a.season_year] = target.get(a.season_year, 0.0) + a.kg
    for y, kg in annual_block.items():
        if y not in totals:
            totals[y] = kg
            per_block_years.add(y)
    for y, kg in annual_farm.items():
        totals.setdefault(y, kg)
    return totals, per_block_years


def _farm_history(totals: dict, season_year: int, current_year: int) -> list:
    """_farm_totals() before `season_year`, oldest first, for the chart."""
    return [{"year": y, "kg": _r(kg), "partial": y == current_year}
            for y, kg in sorted(totals.items()) if y < season_year]


def _progress(season: dict, season_year: int, estimate_kg: Optional[float], today: date) -> Optional[dict]:
    """How the current season is tracking: kg picked so far, what the
    estimate implies should have been picked by now, and a projection of the
    season total from the typical share picked by this day. None for any
    season but the current one."""
    if season_year != season["current_year"]:
        return None
    am, ad, day_kg = season["anchor_month"], season["anchor_day"], season["day_kg"]
    today_sd = season_day(today, season_year, am, ad)

    actual = sum(kg for (y, _, _), kg in day_kg.items() if y == season_year)

    # Each daily-tracked past season: the share of its total picked by the
    # same day of the season.
    year_total: dict = {}
    year_to_date: dict = {}
    for (y, _, d), kg in day_kg.items():
        if y == season_year:
            continue
        year_total[y] = year_total.get(y, 0.0) + kg
        if season_day(d, y, am, ad) <= today_sd:
            year_to_date[y] = year_to_date.get(y, 0.0) + kg
    shares = [{"year": y, "share": round(year_to_date.get(y, 0.0) / t, 3)}
              for y, t in sorted(year_total.items()) if t > 0]
    typical = _mean([s["share"] for s in shares])

    projected = low = high = None
    range_years = []
    if typical is not None and typical >= MIN_PROJECTION_SHARE and actual > 0:
        projected = actual / typical
        usable = [s for s in shares if s["share"] >= MIN_RANGE_SHARE]
        if usable:
            low = actual / max(s["share"] for s in usable)
            high = actual / min(s["share"] for s in usable)
            range_years = [s["year"] for s in usable]

    return {
        "season_day": today_sd,
        "actual_kg": _r(actual),
        "typical_share": typical if typical is None else round(typical, 3),
        "shares": shares,
        "expected_by_now_kg": _r(estimate_kg * typical) if estimate_kg and typical is not None else None,
        "projected_kg": _r(projected),
        "projected_low_kg": _r(low),
        "projected_high_kg": _r(high),
        # The seasons the range is drawn from; the rest had picked less than
        # min_range_share by this day and are left out of it.
        "range_years": range_years,
        "min_projection_share": MIN_PROJECTION_SHARE,
        "min_range_share": MIN_RANGE_SHARE,
    }


# --------------------------------------------------------------------------- #
# Pack-out
# --------------------------------------------------------------------------- #
# Shares may leave a little over 100 from rounding in a hand-typed mix, never
# more: anything past this is two lines counting the same fruit.
_PACK_SHARE_TOLERANCE = 0.05


def _cartons(x: float) -> int:
    """Whole cartons, half up - JavaScript's Math.round, which the tab's
    live figures use. Python's round() goes half to even, and the screen and
    the saved version would then disagree by a carton on an exact half."""
    return math.floor(x + 0.5)


def _channel_key(name: str) -> str:
    """How pack-out lines are grouped and checked for duplicates. lower(),
    not casefold(), to match the tab's toLowerCase()."""
    return name.strip().lower()


def _packout(total_kg: float, rows: list) -> Optional[dict]:
    """Season kg and cartons per pack-out line and per channel, from the
    estimate's net picked kg. The farm's "%" sheet sum: kg = total x share,
    cartons = kg / kg per carton. A line with no kg per carton (juice,
    rejects) carries kg only. Whatever share the lines leave unallocated is
    reported, never spread over them.

    Mirrored by packout() in frontend/shared/estimate-tab.js so the tab can
    follow unsaved edits - keep the two identical."""
    if not rows:
        return None
    total = total_kg or 0.0
    lines_out = []
    channels: dict = {}   # _channel_key -> totals, in first-seen order
    packed_kg = not_packed_kg = cartons_sum = allocated = 0.0
    for r in sorted(rows, key=lambda r: r.position):
        kg = total * r.share_pct / 100
        cartons = kg / r.kg_per_carton if r.kg_per_carton else None
        allocated += r.share_pct
        if cartons is None:
            not_packed_kg += kg
        else:
            packed_kg += kg
            cartons_sum += cartons
        key = _channel_key(r.channel)
        ch = channels.setdefault(key, {"channel": r.channel.strip(), "share_pct": 0.0, "kg": 0.0, "cartons": None})
        ch["share_pct"] += r.share_pct
        ch["kg"] += kg
        if cartons is not None:
            ch["cartons"] = (ch["cartons"] or 0.0) + cartons
        lines_out.append({"position": r.position, "channel": r.channel, "pack_type": r.pack_type,
                          "kg_per_carton": r.kg_per_carton, "share_pct": r.share_pct, "note": r.note,
                          "kg": _r(kg), "cartons": _cartons(cartons) if cartons is not None else None})
    unallocated = 100.0 - allocated
    return {
        "basis_kg": _r(total),
        "lines": lines_out,
        "channels": [{"channel": c["channel"], "share_pct": round(c["share_pct"], 2), "kg": _r(c["kg"]),
                      "cartons": _cartons(c["cartons"]) if c["cartons"] is not None else None}
                     for c in channels.values()],
        "packed_kg": _r(packed_kg),
        "not_packed_kg": _r(not_packed_kg),
        "cartons": _cartons(cartons_sum),
        "avg_kg_per_carton": round(packed_kg / cartons_sum, 2) if cartons_sum else None,
        "allocated_pct": round(allocated, 2),
        "unallocated_pct": round(unallocated, 2),
        "unallocated_kg": _r(total * unallocated / 100),
    }


def _pack_rows(owner: Session, estimate_id: int) -> list:
    return owner.exec(select(YieldEstimatePack).where(YieldEstimatePack.estimate_id == estimate_id)
                      .order_by(YieldEstimatePack.position)).all()


# --------------------------------------------------------------------------- #
# Weather-model cross-check
# --------------------------------------------------------------------------- #
def _crosscheck(total_kg, favorable_kg, expected_kg, unfavorable_kg) -> Optional[dict]:
    """Where an estimate sits against the weather model's three scenarios.
    The Unfavorable-Favorable span is the model's worst- and best-weather
    cases, not an error band, so "within" means no more than that.

    Mirrored by crosscheck() in frontend/shared/estimate-tab.js - keep the
    two identical."""
    if not total_kg or None in (favorable_kg, expected_kg, unfavorable_kg):
        return None
    lo = min(favorable_kg, expected_kg, unfavorable_kg)
    hi = max(favorable_kg, expected_kg, unfavorable_kg)
    return {
        "gap_kg": _r(total_kg - expected_kg),
        # Half up to one decimal, as the tab's Math.round(x * 10) / 10 does.
        "gap_pct": math.floor((total_kg - expected_kg) / expected_kg * 1000 + 0.5) / 10 if expected_kg else None,
        "position": "below" if total_kg < lo else "above" if total_kg > hi else "within",
    }


def _snapshot_out(est: YieldEstimate, total_kg: float) -> Optional[dict]:
    if est.forecast_expected_kg is None or est.forecast_built_at is None:
        return None
    return {
        "built_at": est.forecast_built_at.isoformat() + "Z",
        "favorable_kg": est.forecast_favorable_kg,
        "expected_kg": est.forecast_expected_kg,
        "unfavorable_kg": est.forecast_unfavorable_kg,
        "live": est.forecast_live,
        "settled": est.forecast_settled,
        **(_crosscheck(total_kg, est.forecast_favorable_kg, est.forecast_expected_kg,
                       est.forecast_unfavorable_kg) or {"gap_kg": None, "gap_pct": None, "position": None}),
    }


# --------------------------------------------------------------------------- #
# Estimates
# --------------------------------------------------------------------------- #
def _lines_total(lines: list) -> float:
    return sum(l.trees * l.kg_per_tree for l in lines if l.kg_per_tree is not None)


def _estimate_out(est: YieldEstimate, lines: list, blocks: dict, pack: list) -> dict:
    lines = sorted(lines, key=lambda l: _block_sort_key(l.block_id))
    out_lines = []
    total_trees = 0
    for l in lines:
        kg = l.trees * l.kg_per_tree if l.kg_per_tree is not None else None
        total_trees += l.trees or 0
        out_lines.append({"block_id": l.block_id, "trees": l.trees, "kg_per_tree": l.kg_per_tree,
                          "kg": _r(kg), "note": l.note, "in_register": l.block_id in blocks})
    total_kg = _lines_total(lines)
    return {
        "id": est.id, "season_year": est.season_year, "name": est.name, "notes": est.notes,
        "created_at": est.created_at.isoformat() + "Z", "updated_at": est.updated_at.isoformat() + "Z",
        "lines": out_lines, "total_kg": _r(total_kg), "total_trees": total_trees,
        "blocks_estimated": sum(1 for l in lines if l.kg_per_tree is not None),
        "pack": [{"position": p.position, "channel": p.channel, "pack_type": p.pack_type,
                  "kg_per_carton": p.kg_per_carton, "share_pct": p.share_pct, "note": p.note}
                 for p in pack],
        "packout": _packout(total_kg, pack),
        "forecast_snapshot": _snapshot_out(est, total_kg),
    }


def _lines(owner: Session, estimate_id: int) -> list:
    return owner.exec(select(YieldEstimateBlock).where(YieldEstimateBlock.estimate_id == estimate_id)).all()


def _full_out(owner: Session, est: YieldEstimate, blocks: dict) -> dict:
    return _estimate_out(est, _lines(owner, est.id), blocks, _pack_rows(owner, est.id))


def _get_estimate(owner: Session, estimate_id: int) -> YieldEstimate:
    est = owner.get(YieldEstimate, estimate_id)
    if not est:
        raise HTTPException(404, "No such estimate")
    return est


def _check_season(season_year: int) -> None:
    if not _SEASON_MIN <= season_year <= _SEASON_MAX:
        raise HTTPException(400, f"Season must be between {_SEASON_MIN} and {_SEASON_MAX}")


@router.get("")
def estimate_view(season: Optional[int] = None, estimate_id: Optional[int] = None,
                  boord: Session = Depends(get_boord_session),
                  owner: Session = Depends(get_owner_session)):
    """Everything the Estimate tab shows for one season: each own block's
    reference figures, the farm's season totals, the season's saved
    estimates (each with its total and weather-model snapshot), the chosen
    one in full (the most recently changed unless `estimate_id` says
    otherwise) and, for the current season, progress."""
    s = season_day_kg(boord, owner)
    current_year = s["current_year"]
    season_year = season if season is not None else current_year
    _check_season(season_year)
    blocks = s["blocks"]
    block_year_kg, annual_only = block_season_kg(owner, s)
    totals, _ = _farm_totals(owner, s)

    actual_by_block: dict = {}
    if season_year == current_year:
        for (y, bid, _), kg in s["day_kg"].items():
            if y == season_year and bid:
                actual_by_block[bid] = actual_by_block.get(bid, 0.0) + kg

    ref_blocks = []
    for bid in sorted(blocks, key=_block_sort_key):
        b = blocks[bid]
        if not b.active and bid not in block_year_kg:
            continue
        ref = _block_reference(block_year_kg.get(bid, {}), b.trees or 0, season_year,
                               annual_only, current_year)
        ref_blocks.append({"block_id": bid, "name": b.name, "variety": b.variety,
                           "trees": b.trees, "hectares": b.hectares, "active": b.active,
                           "actual_kg": _r(actual_by_block.get(bid)) if season_year == current_year else None,
                           **ref})

    estimates = owner.exec(select(YieldEstimate).where(YieldEstimate.season_year == season_year)
                           .order_by(YieldEstimate.updated_at.desc(), YieldEstimate.id.desc())).all()
    chosen = None
    if estimate_id is not None:
        chosen = next((e for e in estimates if e.id == estimate_id), None)
        if chosen is None:
            raise HTTPException(404, "No such estimate for that season")
    elif estimates:
        chosen = estimates[0]
    chosen_out = _full_out(owner, chosen, blocks) if chosen else None

    # Every version's total, for the "estimate vs the model at the time vs
    # actual" table - one query for all their lines, not one per version.
    lines_by_est: dict = {}
    if estimates:
        for l in owner.exec(select(YieldEstimateBlock).where(
                YieldEstimateBlock.estimate_id.in_([e.id for e in estimates]))).all():
            lines_by_est.setdefault(l.estimate_id, []).append(l)
    versions = []
    for e in estimates:
        total = _lines_total(lines_by_est.get(e.id, []))
        versions.append({"id": e.id, "name": e.name, "updated_at": e.updated_at.isoformat() + "Z",
                         "total_kg": _r(total), "forecast": _snapshot_out(e, total)})

    season_kg = totals.get(season_year)
    return {
        "season_year": season_year,
        "current_year": current_year,
        "history_years": sorted({y for r in ref_blocks for y in r["history"]}),
        "blocks": ref_blocks,
        "farm_history": _farm_history(totals, season_year, current_year),
        # What the season being viewed actually came to (so far, for the
        # current one) - what every version is judged against afterwards.
        "season_total": ({"kg": _r(season_kg), "partial": season_year == current_year}
                         if season_kg is not None else None),
        "estimates": versions,
        "estimate": chosen_out,
        "progress": _progress(s, season_year, chosen_out["total_kg"] if chosen_out else None,
                              date.today()),
    }


class EstimateLineIn(BaseModel):
    block_id: str = PydField(min_length=1, max_length=40)
    trees: int = PydField(ge=0, le=1_000_000)
    kg_per_tree: Optional[float] = PydField(default=None, ge=0, le=2000)
    note: str = PydField(default="", max_length=500)


class PackLineIn(BaseModel):
    channel: str = PydField(min_length=1, max_length=60)
    pack_type: str = PydField(default="", max_length=60)
    kg_per_carton: Optional[float] = PydField(default=None, gt=0, le=50)
    share_pct: float = PydField(ge=0, le=100)
    note: str = PydField(default="", max_length=200)


class EstimateForecastIn(BaseModel):
    """The Harvest Forecast figures the browser was showing, as built by
    risk.build_harvest_forecast() at `built_at`. See YieldEstimate."""
    season_year: int
    built_at: datetime
    favorable_kg: float = PydField(ge=0, le=1e8)
    expected_kg: float = PydField(ge=0, le=1e8)
    unfavorable_kg: float = PydField(ge=0, le=1e8)
    live: bool
    settled: int = PydField(ge=0, le=10)


class EstimateCreate(BaseModel):
    season_year: int
    name: str = PydField(default="", max_length=100)
    notes: str = PydField(default="", max_length=4000)
    copy_from_id: Optional[int] = None
    forecast: Optional[EstimateForecastIn] = None


class EstimateUpdate(BaseModel):
    name: Optional[str] = PydField(default=None, max_length=100)
    notes: Optional[str] = PydField(default=None, max_length=4000)
    lines: Optional[List[EstimateLineIn]] = None
    pack: Optional[List[PackLineIn]] = None
    forecast: Optional[EstimateForecastIn] = None


def _replace_lines(owner: Session, estimate_id: int, lines: list) -> None:
    seen = set()
    for l in lines:
        bid = l.block_id.strip()
        if not bid:
            raise HTTPException(400, "A line has no block")
        if bid in seen:
            raise HTTPException(400, f"Block {bid} appears twice")
        seen.add(bid)
    owner.exec(delete(YieldEstimateBlock).where(YieldEstimateBlock.estimate_id == estimate_id))
    owner.add_all(YieldEstimateBlock(estimate_id=estimate_id, block_id=l.block_id.strip(),
                                     trees=l.trees, kg_per_tree=l.kg_per_tree, note=l.note.strip())
                  for l in lines)


def _replace_pack(owner: Session, estimate_id: int, pack: list) -> None:
    seen = set()
    total = 0.0
    for p in pack:
        channel = p.channel.strip()
        if not channel:
            raise HTTPException(400, "Every pack-out line needs a channel")
        key = (_channel_key(channel), _channel_key(p.pack_type))
        if key in seen:
            label = f"{channel} {p.pack_type.strip()}".strip()
            raise HTTPException(400, f"Pack-out line {label} appears twice")
        seen.add(key)
        total += p.share_pct
    if total > 100 + _PACK_SHARE_TOLERANCE:
        raise HTTPException(400, f"Pack-out shares add up to {total:.1f}% - more than all the fruit")
    owner.exec(delete(YieldEstimatePack).where(YieldEstimatePack.estimate_id == estimate_id))
    owner.add_all(YieldEstimatePack(estimate_id=estimate_id, position=i, channel=p.channel.strip(),
                                    pack_type=p.pack_type.strip(), kg_per_carton=p.kg_per_carton,
                                    share_pct=p.share_pct, note=p.note.strip())
                  for i, p in enumerate(pack))


def _naive_utc(dt: datetime) -> datetime:
    """A client timestamp as this app stores them: naive UTC. One without a
    zone is taken as UTC already - the server sends them with a "Z"."""
    return dt.astimezone(timezone.utc).replace(tzinfo=None) if dt.tzinfo else dt


def _apply_forecast(est: YieldEstimate, forecast: EstimateForecastIn, current_year: int) -> None:
    """Stores a Harvest Forecast snapshot on a version, refusing one that
    cannot be what it claims: the model only ever describes the current
    season, and a forecast "built" in the future was never built."""
    if not (forecast.season_year == est.season_year == current_year):
        raise HTTPException(400, f"The weather model only covers the current season ({current_year})")
    built_at = _naive_utc(forecast.built_at)
    if built_at > _utcnow() + timedelta(minutes=10):
        raise HTTPException(400, "That forecast's time is in the future")
    est.forecast_favorable_kg = forecast.favorable_kg
    est.forecast_expected_kg = forecast.expected_kg
    est.forecast_unfavorable_kg = forecast.unfavorable_kg
    est.forecast_built_at = built_at
    est.forecast_live = forecast.live
    est.forecast_settled = forecast.settled


def _latest_pack_source(owner: Session) -> Optional[YieldEstimate]:
    """The most recently changed version, of any season, that has a pack-out
    mix - where a brand-new estimate takes its starting mix from."""
    return owner.exec(select(YieldEstimate)
                      .where(YieldEstimate.id.in_(select(YieldEstimatePack.estimate_id)))
                      .order_by(YieldEstimate.updated_at.desc(), YieldEstimate.id.desc())).first()


def _register(boord: Session) -> tuple:
    """The current season and this farm's own blocks, straight from Boord's
    Settings and block register - all a save needs. season_day_kg() would
    also scan every crate picked this season, which a save has no use for."""
    settings = boord.exec(select(SystemSetting)).first()
    current_year = settings.current_harvest_year if settings else date.today().year
    own = own_farm_block_ids(boord)
    return current_year, {b.id: b for b in boord.exec(select(Block)).all() if b.id in own}


@router.post("")
def create_estimate(body: EstimateCreate, boord: Session = Depends(get_boord_session),
                    owner: Session = Depends(get_owner_session)):
    """A new version for a season. Copies another version's block lines and
    pack-out mix when `copy_from_id` is given (the usual case: revise the
    last estimate); otherwise starts from Boord's block register - every
    own, active block with its current tree count and no kg/tree yet - and
    takes its pack-out mix from the latest version that has one, in any
    season (`pack_copied_from` says which). A weather-model snapshot is
    never copied: it belongs to the moment its own version was saved."""
    _check_season(body.season_year)
    current_year, blocks = _register(boord)
    pack_source = None
    if body.copy_from_id is not None:
        pack_source = _get_estimate(owner, body.copy_from_id)
        src = [EstimateLineIn(block_id=l.block_id, trees=l.trees, kg_per_tree=l.kg_per_tree, note=l.note)
               for l in _lines(owner, body.copy_from_id)]
    else:
        src = [EstimateLineIn(block_id=b.id, trees=b.trees or 0)
               for b in blocks.values() if b.active]
        pack_source = _latest_pack_source(owner)
    pack = [PackLineIn(channel=p.channel, pack_type=p.pack_type, kg_per_carton=p.kg_per_carton,
                       share_pct=p.share_pct, note=p.note)
            for p in (_pack_rows(owner, pack_source.id) if pack_source else [])]

    now = _utcnow()
    est = YieldEstimate(season_year=body.season_year, notes=body.notes.strip(),
                        name=body.name.strip() or date.today().strftime("%d %b %Y"),
                        created_at=now, updated_at=now)
    if body.forecast is not None:
        _apply_forecast(est, body.forecast, current_year)
    owner.add(est)
    owner.flush()
    _replace_lines(owner, est.id, src)
    _replace_pack(owner, est.id, pack)
    owner.commit()
    owner.refresh(est)
    out = _full_out(owner, est, blocks)
    out["pack_copied_from"] = ({"id": pack_source.id, "season_year": pack_source.season_year,
                                "name": pack_source.name} if pack_source and pack else None)
    return out


@router.put("/{estimate_id}")
def update_estimate(estimate_id: int, body: EstimateUpdate, boord: Session = Depends(get_boord_session),
                    owner: Session = Depends(get_owner_session)):
    """Rename, re-note, or replace the block lines or pack-out mix of one
    version, and optionally store the weather-model snapshot the browser was
    showing. `lines` and `pack`, when given, are the whole set - a block or
    pack-out line left out is dropped. Anything not sent is left as it was,
    the snapshot included."""
    est = _get_estimate(owner, estimate_id)
    current_year, blocks = _register(boord)
    if body.name is not None:
        if not body.name.strip():
            raise HTTPException(400, "An estimate needs a name")
        est.name = body.name.strip()
    if body.notes is not None:
        est.notes = body.notes.strip()
    if body.forecast is not None:
        _apply_forecast(est, body.forecast, current_year)
    if body.lines is not None:
        _replace_lines(owner, est.id, body.lines)
    if body.pack is not None:
        _replace_pack(owner, est.id, body.pack)
    est.updated_at = _utcnow()
    owner.add(est)
    owner.commit()
    owner.refresh(est)
    return _full_out(owner, est, blocks)


@router.delete("/{estimate_id}")
def delete_estimate(estimate_id: int, owner: Session = Depends(get_owner_session)):
    est = _get_estimate(owner, estimate_id)
    owner.exec(delete(YieldEstimateBlock).where(YieldEstimateBlock.estimate_id == estimate_id))
    owner.exec(delete(YieldEstimatePack).where(YieldEstimatePack.estimate_id == estimate_id))
    owner.delete(est)
    owner.commit()
    return {"deleted": estimate_id}


@router.get("/{estimate_id}/export")
def export_estimate(estimate_id: int, boord: Session = Depends(get_boord_session),
                    owner: Session = Depends(get_owner_session)):
    """One version as a workbook, laid out like the farm's own Produksie
    sheet: per block, the history it was judged against and the estimate;
    then its pack-out, if it has a mix."""
    est = _get_estimate(owner, estimate_id)
    view = estimate_view(season=est.season_year, estimate_id=est.id, boord=boord, owner=owner)
    e = view["estimate"]
    ref = {b["block_id"]: b for b in view["blocks"]}
    years = view["history_years"][-3:]
    is_current = est.season_year == view["current_year"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estimate"
    ws.append([f"Crop estimate {est.season_year}: {est.name}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Last changed {est.updated_at:%d %b %Y %H:%M} UTC"])
    if est.notes:
        ws.append([est.notes])
    ws.append([])
    header = (["Block", "Variety", "Trees"] + [f"{y} kg/tree" for y in years]
              + ["Last 5 avg kg/tree", "Best 5 of 10 kg/tree", "Estimate kg/tree", "Estimate kg"]
              + (["Picked kg"] if is_current else []) + ["Note"])
    ws.append(header)
    for c in ws[ws.max_row]:
        _style_header_cell(c)
    header_row = ws.max_row
    for l in e["lines"]:
        r = ref.get(l["block_id"], {})
        hist = r.get("history", {})
        ws.append([l["block_id"], r.get("variety", ""), l["trees"]]
                  + [hist.get(y, {}).get("kg_tree") for y in years]
                  + [r.get("avg5_kg_tree"), r.get("best5_kg_tree"), l["kg_per_tree"], l["kg"]]
                  + ([r.get("actual_kg")] if is_current else []) + [l["note"]])
    ws.append(["Total", "", e["total_trees"]] + [None] * len(years) + [None, None, None, e["total_kg"]]
              + ([view["progress"]["actual_kg"]] if is_current and view["progress"] else []))
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions[openpyxl.utils.get_column_letter(len(header))].width = 40

    po = e["packout"]
    if po:
        ps = wb.create_sheet("Pack-out")
        ps.append([f"Pack-out of {po['basis_kg']:,.0f} kg net picked (estimate {est.season_year}: {est.name})"])
        ps["A1"].font = Font(bold=True)
        ps.append([])
        ps.append(["Channel", "Pack type", "Kg per carton", "% of picked kg", "Kg", "Cartons", "Note"])
        for c in ps[ps.max_row]:
            _style_header_cell(c)
        for l in po["lines"]:
            ps.append([l["channel"], l["pack_type"], l["kg_per_carton"], l["share_pct"], l["kg"],
                       l["cartons"], l["note"]])
        ps.append([])
        ps.append(["By channel", "", "", "% of picked kg", "Kg", "Cartons"])
        for c in ps[ps.max_row]:
            c.font = Font(bold=True)
        for c in po["channels"]:
            ps.append([c["channel"], "", "", c["share_pct"], c["kg"], c["cartons"]])
        ps.append(["Total", "", po["avg_kg_per_carton"], po["allocated_pct"], _r(po["packed_kg"] + po["not_packed_kg"]),
                   po["cartons"]])
        for c in ps[ps.max_row]:
            c.font = Font(bold=True)
        if po["unallocated_pct"] > _PACK_SHARE_TOLERANCE:
            ps.append(["Unallocated", "", "", po["unallocated_pct"], po["unallocated_kg"]])
        ps.column_dimensions["A"].width = 20
        ps.column_dimensions["B"].width = 14
        ps.column_dimensions["G"].width = 30

    fs = wb.create_sheet("Farm by season")
    fs.append(["Season", "Kg"])
    for c in fs[1]:
        _style_header_cell(c)
    for h in view["farm_history"]:
        fs.append([f"{h['year']}{' (in progress)' if h['partial'] else ''}", h["kg"]])

    # Nothing in this workbook is meant to be a formula, but openpyxl makes
    # one of any text starting with "=" - a note like "=same as 8a" would
    # open as #NAME? (and run as a formula). Keep every such cell as text.
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    filename = f"Crop_Estimate_{est.season_year}_{est.id}.xlsx"
    # Kept on disk as well, like every workbook this app builds - see
    # routers/historical_report.py's REPORTS_DIR.
    with open(os.path.join(REPORTS_DIR, filename), "wb") as f:
        f.write(data)
    return Response(data, media_type=XLSX_MEDIA,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})
