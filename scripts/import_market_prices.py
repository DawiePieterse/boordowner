#!/usr/bin/env python3
"""Import of the farm's hand-kept litchi "Markte" market-statement workbook,
one per season, into MarketPackPrice (season totals: price per carton by
buyer and pack type) and MarketDailySale (a row per delivery, for the
Markets tab's daily price chart). Nothing is read from a fixed column
layout: a buyer or pack type can be added from one season's workbook to the
next, so the pack-type columns are found by their own SEND/SOLD/
OUTSTANDING/PRICE/SUB headers.

Source: a per-season workbook such as
data/imports/LIETJIES Markte 2025 SEISOEN.xlsx (gitignored - per-farm data,
not shipped with the app). The office has kept these since 2022, each
season a little differently, and all of them are read:

  * 2024 and 2025 have a "<year> Totaal" sheet, which is the season's
    totals, and a sheet per receiver (ToroFruit, Dapper, AgriLink USA
    Lug...) for the daily figures.
  * 2022 and 2023 have the per-receiver sheets only, so their totals are
    added up from those (derived_totals). 2022's sheets are an older layout
    still (see _layout), and its local markets share one "Local" sheet
    (MIXED_SHEETS).

A workbook is the record of ONE season, its own: the newest of its Totaal
sheets, or for a workbook with none the year in its file name. The office
also copies last season's Totaal sheet into the new workbook for
comparison, but those copies are not imported - the 2024 workbook's
"2023 Totaal" has a row of formula junk (Dapper's cartons at R0.03 to
R2.33 a carton) and 2023's 4.5kg air sales filed under "2kg MIX", because
the copy kept 2024's column headings. A season's own workbook is the one to
believe, and import_market_prices.py run on it sets that season right
whatever another workbook said.

Safe to re-run: replaces only the workbook's own season, so importing next
season's workbook never touches an earlier one.

Usage (run with the backend's own venv so sqlmodel etc. are on the path):
    backend/.venv/bin/python3 scripts/import_market_prices.py <path-to-workbook.xlsx>
"""
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime

import openpyxl

BACKEND_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "backend")
sys.path.insert(0, BACKEND_DIR)

from sqlmodel import Session, delete  # noqa: E402

from db import init_owner_db as run_migrations, owner_engine as engine  # noqa: E402
from models_owner import MarketDailySale, MarketPackPrice  # noqa: E402

_SHEET_RE = re.compile(r"^(\d{4})\s+totaal$", re.IGNORECASE)
_FILE_YEAR_RE = re.compile(r"(20\d\d)")

# Pack-type names the workbooks and Boord Packhouse's own Settings ->
# Consignment list never agreed on at all - an Afrikaans word ("Kratte" is
# "crate"), the farm's own nickname for a size ("Rhino" for the 4.5kg XL
# carton), or the same carton written as 4.55kg in 2022. routers/market.py's
# _pack_key() already bridges case and punctuation when matching against
# Packhouse's live list, but it can only ever match the same word
# differently spelled - it has no way to know these are the same pack type
# at all. Matched ignoring capitals ("4.5Kg XL" in 2023). Same shape as
# block_renames.py: a record of this one farm's own naming, not a general
# mechanism. Add to this map rather than editing the workbooks, which are
# the farm's own record and should keep saying what they said when they
# were written.
PACK_RENAMES = {
    "15 Kg Kratte": "15kg Crate",
    "15kg Krat": "15kg Crate",     # 2023, TFFG
    "Kratte": "15kg Crate",        # 2023, Tomahawk - the office's own 2023 total filed it as 15 Kg Kratte
    "4.5kg XL": "4.5kg Rhino",
    "4.55kg XL": "4.5kg Rhino",    # 2022
    "4.55kg XXL": "4.5kg XXL",     # 2022
    "PrePac L": "PrePac",          # 2022: the old layout gave every pack a grade, PrePac too
}

# Per-receiver sheet -> the receiver Boord Packhouse has for it (Settings ->
# Delivery Note), so a workbook season and a live season chart under the
# same names. "Lug" is air freight and "See" sea; the office kept AgriLink's
# 2kg air boxes on a sheet of their own, but they are the same receiver, so
# a day on both becomes one point, its price weighted by cartons
# (routers/market.py _daily). AgriLink's sea sheets are USA in every
# season (2024's rows say so; 2022 and 2023 went on the same LW export
# notes), whatever each season called the sheet. AgriLink's other air
# sheets keep one name of their own: 2024's went to NL/AE, not the USA.
# Only names that differ by more than case or punctuation need an entry:
# routers/market.py matches the rest against Packhouse's live receiver
# list. Same shape and rule as PACK_RENAMES, and also applied to each row's
# MARK on a MIXED_SHEETS sheet.
RECEIVER_RENAMES = {
    "AgriLink 2kg Lug": "Agrilink USA Air",
    "AgriLink USA Lug": "Agrilink USA Air",
    "AgriLink Lug USA": "Agrilink USA Air",   # 2024
    "AgriLink USA See": "Agrilink USA Sea",
    "AgriLink See": "Agrilink USA Sea",       # 2022, 2024
    "Agrilink SV 23": "Agrilink USA Sea",     # 2023
    "Agrilink Lug": "AgriLink Lug",           # 2023's spelling of 2022's and 2024's sheet
    "BoereMT": "Boere Mark Trust",
    "BoereTrust": "Boere Mark Trust",         # 2023
    "Boeretrust": "Boere Mark Trust",         # 2024, and 2022's MARK
    "TFFG See": "TFFGSA",
    "Toro Fruit": "ToroFruit",                # 2022, 2023
}

# Sheets that hold several receivers rather than one: 2022's "Local" kept
# every local market (Dapper, Boeretrust, TFFG...) on one sheet, a row each,
# so each row's receiver is its own MARK column, not the sheet name.
MIXED_SHEETS = {"Local"}

# Per season, sheets that copy deliveries another sheet in the same
# workbook already has (the office's habit of duplicating a sheet before
# reworking it), which would count the same sales twice. Each was compared
# against the others delivery note by delivery note.
#   2022: "AgriLink Lug" (air) and "AgriLink See" (sea) are the season's
#         statements. "AgriLink" is both put together, with the sea notes
#         at the R87.63 advance payment instead of what the market made;
#         "AgriLink (3)" and "AgriLink (5)" are part-copies of it.
#   2023: "Agrilink LU 23" is a part-copy of "Agrilink Lug".
SKIP_SHEETS = {
    2022: {"AgriLink", "AgriLink (3)", "AgriLink (5)"},
    2023: {"Agrilink LU 23"},
}

_PACK_RENAMES = {k.lower(): v for k, v in PACK_RENAMES.items()}


def _clean(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def _pack_type(label_cell, season_year: int) -> str:
    """The group label, with a leading "<season_year> " stripped - one
    season's sheet (2023) wrote its pack-type headers as "2023 2kg XL"
    rather than "2kg XL", which would otherwise stop that season's own
    pack types from matching the same pack type in any other season - and
    passed through PACK_RENAMES."""
    label = _clean(label_cell)
    prefix = f"{season_year} "
    if label.startswith(prefix):
        label = label[len(prefix):]
    return _PACK_RENAMES.get(label.lower(), label)


def _group_starts(header_row: list) -> list:
    """Index of each pack-type column group's SEND cell - wherever the
    SEND/SOLD/OUTSTANDING/PRICE/SUB headers repeat, not a fixed offset, so
    an extra pack-type column doesn't silently shift anything out of
    alignment."""
    return [i for i, v in enumerate(header_row) if v == "SEND"]


def _season_rows(ws, season_year: int) -> list:
    label_row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]   # pack-type name
    header_row = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]  # SEND/SOLD/.../SUB
    groups = [(i, _pack_type(label_row[i], season_year)) for i in _group_starts(header_row)]
    rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        buyer = _clean(row[1]) if len(row) > 1 else ""
        if not buyer or "total" in buyer.lower() or "totaal" in buyer.lower():
            continue
        for start, pack_type in groups:
            if not pack_type:
                continue
            sold = row[start + 1] if start + 1 < len(row) else None
            price = row[start + 3] if start + 3 < len(row) else None
            sub = row[start + 4] if start + 4 < len(row) else None
            if not isinstance(sold, (int, float)) or not sold:
                continue
            if not isinstance(price, (int, float)) or not price:
                continue  # '#DIV/0!' (a formula error) or genuinely 0 - neither is a price
            rows.append(MarketPackPrice(
                season_year=season_year, buyer=buyer, pack_type=pack_type,
                cartons_sold=float(sold), price_per_carton=float(price),
                sub_total=float(sub) if isinstance(sub, (int, float)) else 0.0,
            ))
    return rows


def _sale_date(cell):
    """The Datum column. Each season's office typed it its own way: as text,
    "07.11.25" (2025, 2024) or "13.10.2022" (2022, part of 2023), or as a
    real date cell (most of 2023) - all taken. Anything else is None, never a
    guess: a sale filed under the wrong day would be worse than a missing
    one."""
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, date):
        return cell
    text = _clean(cell)
    for fmt in ("%d.%m.%y", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _layout(ws):
    """Where a per-receiver sheet keeps its figures: (header row index,
    [(first column, pack type label), ...], {column name: index}), or None
    for a sheet with no PRICE header at all (a bank-payment list, a picking
    record, a summary of cartons only - nothing to chart).

    The header row is found by its PRICE cells within the first few rows,
    because the sheets don't all start on the same row ("Andre" in 2025, the
    2022 export receivers a row lower than the rest). A pack-type group is
    the columns up to its PRICE - SUB after it is not required, since some
    2023 sheets have it on the row above instead - and its label is read
    two ways:
      * 2023 onward, "SEND, SOLD, OUTSTANDING, PRICE, SUB" under a label row
        naming the pack type ("2kg XXL");
      * 2022, the group's first column holds the grade instead of "SEND"
        ("XXL, Sold, Outstandin, PRICE, SUB") under a row with only the
        carton size ("4.55kg"), so the pack type is the two put together.
    The column names come from both rows, since some 2023 sheets put
    "Datum, MARK, NR" on the label row rather than the header row."""
    top = [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_row=6)]
    hdr = next((i for i, r in enumerate(top) if any(_clean(v).upper() == "PRICE" for v in r)), None)
    if not hdr:
        return None
    head, label = top[hdr], top[hdr - 1]
    cell = lambda row, i: _clean(row[i]) if 0 <= i < len(row) else ""
    groups = []
    for p in range(3, len(head)):
        if cell(head, p).upper() != "PRICE" or not cell(head, p - 2).upper().startswith("SOLD"):
            continue
        start = p - 3
        first = cell(head, start)
        groups.append((start, cell(label, start) if first.upper() == "SEND"
                       else f"{cell(label, start)} {first}".strip()))
    cols = {}
    for row in (label, head):
        for i, v in enumerate(row):
            cols.setdefault(_clean(v).upper(), i)
    return hdr, groups, cols


def _daily_rows(ws, receiver: str, season_year: int, undated: list, notes: list) -> list:
    """A per-receiver sheet's priced lines, one MarketDailySale per pack
    type per row. The receiver is the sheet's own name, or each row's MARK
    for a sheet that holds several (MIXED_SHEETS).

    A row with no Datum is one of two things. Where the office put a second
    line on the same delivery note (two cultivars, two pallets) it left the
    date out, so the line belongs to the row above it. Where a row adds up
    the sheet - its cartons equal everything above it, pack type by pack
    type - it is the totals row, and is passed over; so is any blank-dated
    row before the first delivery (the 2024 sheets' summary line at the
    top). A row whose Datum is a typo rather than a date ("26/09/202") is
    kept off the chart, since its day is unknown, but its lines are added to
    `undated`, so the season totals still count what it sold, and `notes`
    says which."""
    layout = _layout(ws)
    if not layout:
        return []
    hdr, groups, cols = layout
    groups = [(start, _pack_type(label, season_year)) for start, label in groups]
    mark_col = cols.get("MARK", 1)
    note_col = cols.get("NR", cols.get("VRAG NR", 3))
    mixed = _clean(ws.title) in MIXED_SHEETS
    cell = lambda row, i: _clean(row[i]) if i < len(row) else ""
    running = defaultdict(float)   # pack type -> cartons sold so far on this sheet
    rows, last = [], None          # last: (sale_date or None, receiver, note_no) of the delivery above
    for row in ws.iter_rows(min_row=hdr + 2, values_only=True):
        lines = []
        for start, pack_type in groups:
            sold = row[start + 1] if start + 1 < len(row) else None
            price = row[start + 3] if start + 3 < len(row) else None
            if not pack_type or not isinstance(sold, (int, float)) or not sold:
                continue
            if not isinstance(price, (int, float)) or not price:
                continue  # '#DIV/0!' (a formula error) or genuinely 0 - neither is a price
            lines.append((pack_type, float(sold), float(price)))
        if cell(row, 0):
            who = receiver
            if mixed:
                who = RECEIVER_RENAMES.get(cell(row, mark_col), cell(row, mark_col))
            last = (_sale_date(row[0]), who, cell(row, note_col))
            if last[0] is None and lines:
                notes.append(f"{ws.title}: Datum {cell(row, 0)!r}, NR {last[2]} - in the season totals, "
                             "not on the daily chart")
        elif not lines or last is None or all(abs(sold - running[pack]) < 0.5 for pack, sold, _ in lines):
            continue   # nothing priced, before the first delivery, or the totals row
        for pack, sold, _ in lines:
            running[pack] += sold
        sale_date, who, note_no = last if last else (None, receiver, "")
        if sale_date is None:
            undated += [(who, pack, sold, price) for pack, sold, price in lines]
            continue
        rows += [MarketDailySale(season_year=season_year, sale_date=sale_date, receiver=who, note_no=note_no,
                                 pack_type=pack, cartons_sold=sold, price_per_carton=price)
                 for pack, sold, price in lines]
    return rows


def derived_totals(daily: list, undated: list, season_year: int) -> list:
    """Season totals for a workbook with no "<year> Totaal" sheet (2022's,
    2023's), added up from its own per-receiver lines - the daily ones and
    the `undated` (receiver, pack, cartons, price) ones whose day was a
    typo: one MarketPackPrice per receiver and pack type, its price the
    cartons-weighted average. Only priced lines count, which is all a
    Totaal sheet's PRICE could ever have been worked out from either."""
    acc = defaultdict(lambda: [0.0, 0.0])   # (receiver, pack) -> [cartons, gross]
    for receiver, pack, sold, price in ([(d.receiver, d.pack_type, d.cartons_sold, d.price_per_carton)
                                         for d in daily] + undated):
        a = acc[(receiver, pack)]
        a[0] += sold
        a[1] += sold * price
    return [MarketPackPrice(season_year=season_year, buyer=buyer, pack_type=pack_type,
                            cartons_sold=round(cartons, 2), price_per_carton=round(gross / cartons, 2),
                            sub_total=round(gross, 2))
            for (buyer, pack_type), (cartons, gross) in sorted(acc.items()) if cartons]


def load_rows(xlsx_path: str) -> dict:
    """The workbook's own season, read but not stored:
      season        - the newest "<year> Totaal" sheet's year, else the year
                      in the file name ("LIETJIES Markte 2022 SEISOEN.xlsx"),
                      else None (and nothing else is read: no guessed season);
      totals        - its MarketPackPrice rows: its own Totaal sheet, or
                      added up from the per-receiver sheets without one;
      from_totaal   - which of the two;
      daily         - its MarketDailySale rows from the per-receiver sheets
                      (less the season's SKIP_SHEETS);
      older_totaal  - the years of the earlier seasons' Totaal sheets copied
                      in, which are not imported (see the module docstring);
      notes         - lines worth printing (Datum typos)."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    totaal = {}   # year -> its Totaal sheet's name
    for name in wb.sheetnames:
        m = _SHEET_RE.match(name.strip())
        if m:
            totaal[int(m.group(1))] = name
    season = max(totaal) if totaal else None
    if season is None:
        m = _FILE_YEAR_RE.search(os.path.basename(xlsx_path))
        season = int(m.group(1)) if m else None
    out = {"season": season, "totals": [], "from_totaal": season in totaal, "daily": [],
           "older_totaal": sorted(y for y in totaal if y != season), "notes": []}
    if season is None:
        return out
    skip = SKIP_SHEETS.get(season, set())
    undated = []
    for name in wb.sheetnames:
        if not _SHEET_RE.match(name.strip()) and _clean(name) not in skip:
            receiver = RECEIVER_RENAMES.get(_clean(name), _clean(name))
            out["daily"] += _daily_rows(wb[name], receiver, season, undated, out["notes"])
    out["totals"] = (_season_rows(wb[totaal[season]], season) if out["from_totaal"]
                     else derived_totals(out["daily"], undated, season))
    return out


def main():
    if len(sys.argv) < 2:
        print("Usage: import_market_prices.py <path-to-workbook.xlsx>\n"
              "Reads one season's hand-kept litchi market-statement workbook (the\n"
              "farm's own file, not shipped with the app): its '<year> Totaal'\n"
              "sheet, and a sheet per receiver.")
        return
    xlsx_path = sys.argv[1]
    if not os.path.exists(xlsx_path):
        print(f"No such file: {xlsx_path}")
        return
    got = load_rows(xlsx_path)
    season = got["season"]
    if season is None:
        print(f"No '<year> Totaal' sheet and no year in the file name of {xlsx_path} - "
              "nothing imported, since there is no telling which season it is.")
        return
    receivers = len({d.receiver for d in got["daily"]})   # before the commit expires them
    totals, daily = len(got["totals"]), len(got["daily"])
    run_migrations()
    with Session(engine) as session:
        session.exec(delete(MarketPackPrice).where(MarketPackPrice.season_year == season))
        session.add_all(got["totals"])
        session.exec(delete(MarketDailySale).where(MarketDailySale.season_year == season))
        session.add_all(got["daily"])
        session.commit()
    source = f"its '{season} Totaal' sheet" if got["from_totaal"] else "its per-receiver sheets, added up"
    print(f"Imported the {season} season from:\n  {xlsx_path}\n"
          f"  {totals} season-total rows, from {source}\n"
          f"  {daily} daily sales from {receivers} receivers")
    if got["older_totaal"]:
        print(f"  Not imported: the copied-in {', '.join(str(y) for y in got['older_totaal'])} Totaal "
              "sheet(s). Import each season from its own workbook.")
    for line in got["notes"]:
        print(f"  Datum not a date: {line}")


if __name__ == "__main__":
    main()
