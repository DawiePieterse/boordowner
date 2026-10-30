"""scripts/import_market_prices.py: the farm's hand-kept litchi "Markte"
workbooks, one per season since 2022, into MarketPackPrice and
MarketDailySale. The farm's own workbooks are per-farm data and never in
the repo, so each test builds a small workbook in the layout of the season
it stands for: 2024/2025 (a Totaal sheet and SEND/SOLD/... receiver
sheets), 2023 (no Totaal, dates as date cells, a SUB header on the wrong
row) and 2022 (no Totaal, the grade where SEND would be, one "Local" sheet
for every local market)."""
import os
import sys
from datetime import date, datetime

import openpyxl
from sqlmodel import Session, select

from db import owner_engine
from models_owner import MarketDailySale, MarketPackPrice

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "scripts"))
import import_market_prices as imp  # noqa: E402

GROUP = ["SEND", "SOLD", "OUTSTANDING", "PRICE", "SUB"]


def _sheet(wb, title, label, head, rows, blank_first=True):
    """A sheet: an empty first row (as the office's have), the label row,
    the header row, then the data rows."""
    ws = wb.create_sheet(title)
    if blank_first:
        ws.append([])
    ws.append(label)
    ws.append(head)
    for r in rows:
        ws.append(r)
    return ws


def _receiver_2025(wb, title, rows):
    """A 2024/2025 per-receiver sheet: Datum, MARK, DESTIN, NR, then two
    pack-type groups, 2kg XXL and 2kg XL."""
    label = [None] * 4 + ["2kg XXL", None, None, None, None, "2kg XL", None, None, None, None]
    head = ["Datum", "MARK", "DESTIN", "NR"] + GROUP + GROUP
    return _sheet(wb, title, label, head, rows)


def _totaal(wb, title, rows):
    """A '<year> Totaal' sheet: buyer in column B, the same two groups."""
    label = [None] * 4 + ["2kg XXL", None, None, None, None, "2kg XL", None, None, None, None]
    head = ["Datum", "MARK", "DESTIN", "NR"] + GROUP + GROUP
    return _sheet(wb, title, label, head, rows)


def _save(wb, tmp_path, name):
    del wb["Sheet"]
    path = str(tmp_path / name)
    wb.save(path)
    return path


def _workbook_2025(tmp_path):
    wb = openpyxl.Workbook()
    _receiver_2025(wb, "Dapper", [
        [None, "Dapper", None, None, 30, 30, 0, 200, 6000, 0, 0, 0, None, 0],   # summary line above the first delivery
        ["07.11.25", "Dapper", "LC", 1201, 10, 10, 0, 200, 2000, 0, 0, 0, "#DIV/0!", 0],
        [None, None, None, None, 20, 20, 0, 200, 4000, 0, 0, 0, "#DIV/0!", 0],    # a second line on note 1201
        ["10.11.25", "Dapper", "LC", 1205, 0, 0, 0, "#DIV/0!", 0, 5, 5, 0, 150, 750],
        [None, None, None, None, 30, 30, 0, 200, 6000, 5, 5, 0, 150, 750],        # the totals row
    ])
    _receiver_2025(wb, "AgriLink USA Lug", [
        ["08.11.25", "AgriLink", "USA", "LW3001", 100, 100, 0, 300, 30000, 0, 0, 0, "#DIV/0!", 0],
    ])
    _totaal(wb, "2025 Totaal", [
        [None, "Dapper", None, None, 30, 30, 0, 200, 6000, 5, 5, 0, 150, 750],
        [None, "AgriLink USA Lug", None, None, 100, 100, 0, 300, 30000, 0, 0, 0, "#DIV/0!", 0],
        [None, "TOTAAL", None, None, 130, 130, 0, 0, 36000, 5, 5, 0, 0, 750],
    ])
    _totaal(wb, "2024 Totaal", [   # last season's, copied in for comparison
        [None, "Dapper 2024", None, None, 1, 1, 0, 0.09, 0.09, 0, 0, 0, "#DIV/0!", 0],
    ])
    return _save(wb, tmp_path, "LIETJIES Markte 2025 SEISOEN.xlsx")


def test_a_workbook_is_its_own_newest_totaal_season(tmp_path):
    got = imp.load_rows(_workbook_2025(tmp_path))
    assert got["season"] == 2025 and got["from_totaal"] and got["older_totaal"] == [2024]
    assert sorted((t.buyer, t.pack_type, t.cartons_sold, t.price_per_carton) for t in got["totals"]) == [
        ("AgriLink USA Lug", "2kg XXL", 100, 300), ("Dapper", "2kg XL", 5, 150), ("Dapper", "2kg XXL", 30, 200)]


def test_a_second_line_joins_the_delivery_above_and_the_totals_row_is_skipped(tmp_path):
    """Note 1201's undated second line is a sale on 7 November; the summary
    line above the first delivery and the totals row at the bottom are
    neither, or Dapper's 2kg XXL would be counted two or three times."""
    got = imp.load_rows(_workbook_2025(tmp_path))
    dapper = sorted((d.sale_date, d.note_no, d.pack_type, d.cartons_sold)
                    for d in got["daily"] if d.receiver == "Dapper")
    assert dapper == [(date(2025, 11, 7), "1201", "2kg XXL", 10), (date(2025, 11, 7), "1201", "2kg XXL", 20),
                      (date(2025, 11, 10), "1205", "2kg XL", 5)]
    assert {d.receiver for d in got["daily"]} == {"Dapper", "Agrilink USA Air"}   # RECEIVER_RENAMES


def test_import_replaces_only_its_own_season_and_ignores_copied_totaal_sheets(client, tmp_path, monkeypatch, capsys):
    with Session(owner_engine) as owner:
        owner.add_all([
            MarketPackPrice(season_year=2024, buyer="Dapper", pack_type="PrePac", cartons_sold=50,
                            price_per_carton=350, sub_total=17500),
            MarketPackPrice(season_year=2025, buyer="Old", pack_type="2kg XL", cartons_sold=1,
                            price_per_carton=1, sub_total=1),
        ])
        owner.commit()
    monkeypatch.setattr(sys, "argv", ["import_market_prices.py", _workbook_2025(tmp_path)])
    imp.main()
    with Session(owner_engine) as owner:
        totals = owner.exec(select(MarketPackPrice)).all()
        daily = owner.exec(select(MarketDailySale)).all()
    assert {(t.season_year, t.buyer) for t in totals} == {
        (2024, "Dapper"), (2025, "Dapper"), (2025, "AgriLink USA Lug")}   # 2024 untouched by the copied sheet
    assert {d.season_year for d in daily} == {2025} and len(daily) == 4
    assert "Not imported: the copied-in 2024 Totaal" in capsys.readouterr().out


def _workbook_2023(tmp_path):
    """2023: no Totaal sheet, real date cells and a text one, a typo date,
    the 2kg XL group's SUB header on the label row instead of the header
    row, an Afrikaans crate name, and a sheet that copies another."""
    wb = openpyxl.Workbook()
    label = ["Datum", "MARK", "DESTIN", "NR", "2kg XXL", None, None, None, None, "2kg XL", None, None, None, "SUB"]
    head = [None, None, None, None] + GROUP + ["SEND", "SOLD", "OUTSTANDING", "PRICE", None]
    rows = [
        [datetime(2023, 9, 20), "Agrilink", "LU", 919, 65, 65, 0, 345, 22425, 48, 47, 1, 345, 16215],
        ["13.09.2023", "Agrilinl", "LU", 917, 10, 10, 0, 300, 3000, 0, 0, 0, "#DIV/0!", 0],
        ["26/09/202", "Agrilink", "LU", 922, 5, 5, 0, 310, 1550, 0, 0, 0, "#DIV/0!", 0],
    ]
    _sheet(wb, "Agrilink Lug", label, head, rows)
    _sheet(wb, "Agrilink LU 23", label, head, rows[:1])   # a part-copy (SKIP_SHEETS)
    crate_label = [None] * 4 + ["15kg Krat ", None, None, None, None]
    _sheet(wb, "TFFG", crate_label, ["Datum", "MARK", "DESTIN", "NR"] + GROUP,
           [[datetime(2023, 11, 21), "NuLeaf Brands (TFFG)", "LC", 966, 45, 45, 0, 975, 43875]])
    return _save(wb, tmp_path, "LIETJIES Markte 2023 SEISOEN.xlsx")


def test_a_workbook_without_totaal_takes_its_season_from_the_file_name(tmp_path):
    got = imp.load_rows(_workbook_2023(tmp_path))
    assert got["season"] == 2023 and not got["from_totaal"] and got["older_totaal"] == []
    daily = sorted((d.receiver, d.sale_date, d.pack_type, d.cartons_sold) for d in got["daily"])
    assert daily == [
        ("AgriLink Lug", date(2023, 9, 13), "2kg XXL", 10),
        ("AgriLink Lug", date(2023, 9, 20), "2kg XL", 47),    # read without its SUB header
        ("AgriLink Lug", date(2023, 9, 20), "2kg XXL", 65),
        ("TFFG", date(2023, 11, 21), "15kg Crate", 45),       # PACK_RENAMES
    ]   # "Agrilink LU 23" left out: SKIP_SHEETS


def test_a_typo_date_counts_in_the_totals_but_not_on_the_chart(tmp_path):
    got = imp.load_rows(_workbook_2023(tmp_path))
    assert not any(d.note_no == "922" for d in got["daily"])
    xxl = next(t for t in got["totals"] if t.buyer == "AgriLink Lug" and t.pack_type == "2kg XXL")
    # 65 @ 345, 10 @ 300 and note 922's 5 @ 310, weighted by cartons
    assert xxl.cartons_sold == 80 and xxl.sub_total == 22425 + 3000 + 1550
    assert xxl.price_per_carton == round((22425 + 3000 + 1550) / 80, 2)
    assert got["notes"] == ["Agrilink Lug: Datum '26/09/202', NR 922 - in the season totals, not on the daily chart"]


def _workbook_2022(tmp_path):
    """2022's older layout: the header a row lower on some sheets, the
    group's first column the grade rather than SEND, the size alone on the
    label row - and every local market on one "Local" sheet."""
    wb = openpyxl.Workbook()
    label = [None] * 5 + ["4.55kg", None, None, None, None, "PrePac", None, None, None, None]
    head = [None, "MARK", "NR", "WEEK", "KULTIVAR", "XL", "Sold", "Outstandin", "PRICE", "SUB",
            "L", "Sold", "Outstandin", "PRICE", "SUB"]
    ws = _sheet(wb, "AgriLink See", label, head, [
        ["17.11.2022", "AgriLink", "LW2224", 46, "MAU", 4100, 4100, 0, 338, 1385800, 0, 0, 0, "#DIV/0!", 0],
        ["13.10.2022", "AgriLink", "LW2188", 41, "TMR", 0, 0, 0, "#DIV/0!", 0, 384, 384, 0, 284, 109056],
    ])
    ws.insert_rows(1)   # this one starts a row lower
    _sheet(wb, "AgriLink", label, head, [   # the combined copy (SKIP_SHEETS)
        ["17.11.2022", "AgriLink", "LW2224", 46, "MAU", 4100, 4100, 0, 87.65, 359365, 0, 0, 0, "#DIV/0!", 0],
    ])
    local_label = [None] * 6 + ["PrePac", None, None, None, None]
    _sheet(wb, "Local", local_label, [None, "MARK", "DESTIN", "NR", "WEEK", "KULTIVAR"] + GROUP, [
        ["22.09.2022", "Dapper", "LC", 801, 38, "TMR", 3, 3, 0, 320, 960],
        ["12.10.2022", "Boeretrust", "LC", 812, 41, "TMR", 20, 20, 0, 200, 4000],
    ])
    return _save(wb, tmp_path, "LIETJIES Markte 2022 SEISOEN.xlsx")


def test_the_2022_layout_puts_size_and_grade_together(tmp_path):
    got = imp.load_rows(_workbook_2022(tmp_path))
    daily = sorted((d.receiver, d.note_no, d.pack_type, d.cartons_sold, d.price_per_carton) for d in got["daily"])
    assert daily == [
        ("Agrilink USA Sea", "LW2188", "PrePac", 384, 284),        # "PrePac L" -> "PrePac"
        ("Agrilink USA Sea", "LW2224", "4.5kg Rhino", 4100, 338),  # "4.55kg XL" -> "4.5kg Rhino"
        ("Boere Mark Trust", "812", "PrePac", 20, 200),           # Local: the receiver is the row's MARK
        ("Dapper", "801", "PrePac", 3, 320),
    ]   # and nothing from "AgriLink", the combined copy at the advance-payment price


def test_season_totals_are_added_up_per_receiver_and_pack(client, tmp_path, monkeypatch):
    monkeypatch.setattr(sys, "argv", ["import_market_prices.py", _workbook_2022(tmp_path)])
    imp.main()
    r = client.get("/api/market/summary").json()
    assert r["seasons"] == [2022]
    # PrePac: 384 @ 284 + 3 @ 320 + 20 @ 200, weighted by cartons
    assert r["by_season"]["2022"]["PrePac"] == {"cartons": 407.0,
                                                "rand_per_carton": round((109056 + 960 + 4000) / 407, 2)}
    assert r["by_season"]["2022"]["4.5kg Rhino"] == {"cartons": 4100.0, "rand_per_carton": 338.0}
    assert r["daily"]["workbook_seasons"] == [2022]


def test_no_season_no_import(client, tmp_path, monkeypatch, capsys):
    wb = openpyxl.Workbook()
    _receiver_2025(wb, "Dapper", [["07.11.25", "Dapper", "LC", 1201, 10, 10, 0, 200, 2000, 0, 0, 0, None, 0]])
    path = _save(wb, tmp_path, "Markte.xlsx")
    assert imp.load_rows(path)["season"] is None
    monkeypatch.setattr(sys, "argv", ["import_market_prices.py", path])
    imp.main()
    assert "nothing imported" in capsys.readouterr().out
    with Session(owner_engine) as owner:
        assert owner.exec(select(MarketDailySale)).all() == []
