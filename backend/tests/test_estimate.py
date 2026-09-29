"""The Estimate tab: per-block reference figures, saved estimate versions,
and the in-season progress/projection."""
from datetime import date

from db import owner_engine
from models_owner import HistoricalAnnualYield, HistoricalHarvest
from routers.estimate import _block_reference, _progress
from sqlmodel import Session


def _seed_history():
    """Block 7 (2000 trees): seasons 2020-2025 daily, 2016-2019 annual-only.
    Block 8a (500 trees): 2024-2025 daily. Plus one whole-farm season."""
    with Session(owner_engine) as s:
        for year, kg in [(2020, 100_000), (2021, 40_000), (2022, 120_000),
                         (2023, 80_000), (2024, 20_000), (2025, 60_000)]:
            # Two picking days per season, 1 Sep and 1 Nov, split evenly -
            # half the season picked by 1 Sep's season day.
            s.add(HistoricalHarvest(block_id="7", harvest_date=date(year, 9, 1), season_year=year, kg=kg / 2))
            s.add(HistoricalHarvest(block_id="7", harvest_date=date(year, 11, 1), season_year=year, kg=kg / 2))
        for year, kg in [(2016, 200_000), (2017, 160_000), (2018, 60_000), (2019, 90_000)]:
            s.add(HistoricalAnnualYield(block_id="7", season_year=year, kg=kg))
        s.add(HistoricalHarvest(block_id="8a", harvest_date=date(2024, 9, 1), season_year=2024, kg=5_000))
        s.add(HistoricalHarvest(block_id="8a", harvest_date=date(2025, 9, 1), season_year=2025, kg=10_000))
        s.add(HistoricalAnnualYield(block_id=None, season_year=1998, kg=1_148_028))
        s.commit()


def test_reference_figures(client):
    _seed_history()
    body = client.get("/api/estimate?season=2026").json()
    assert body["season_year"] == 2026 and body["current_year"] == 2026
    b7 = next(b for b in body["blocks"] if b["block_id"] == "7")
    # kg/tree on the register's 2000 trees
    assert b7["last_kg_tree"] == 30.0                       # 2025: 60,000 / 2000
    assert b7["avg5_kg_tree"] == round((20 + 60 + 40 + 10 + 30) / 5, 1)   # 2021-2025
    # best 5 of 2016-2025: 100, 80, 60, 50, 45
    assert b7["best5_kg_tree"] == round((100 + 80 + 60 + 50 + 45) / 5, 1)
    assert b7["low_kg_tree"] == 10.0 and b7["high_kg_tree"] == 100.0
    assert b7["history"]["2019"]["annual_only"] is True
    assert b7["history"]["2025"]["annual_only"] is False
    # The current season is never its own reference.
    assert "2026" not in b7["history"]
    # Picked so far this season: the same net kg the Analysis tab counts
    # (the fixture's three crates, plus any a test elsewhere added to the
    # shared fake Boord database).
    analysis = client.get("/api/analysis/summary").json()
    a7 = next(b for b in analysis["block_yield"] if b["block_id"] == "7")
    assert b7["actual_kg"] == a7["by_year"]["2026"]["kg"] >= 51.0
    # Another grower's block never appears.
    assert all(b["block_id"] != "90" for b in body["blocks"])
    # Whole-farm-only seasons still show in the farm totals.
    farm = {h["year"]: h["kg"] for h in body["farm_history"]}
    assert farm[1998] == 1_148_028 and farm[2025] == 70_000
    assert body["estimate"] is None and body["estimates"] == []


def test_block_reference_excludes_unfinished_and_empty_seasons():
    by_year = {2023: 1000.0, 2024: 0.0, 2025: 3000.0}
    ref = _block_reference(by_year, 100, 2025, set(), current_year=2025)
    # Estimating the season being picked: 2025 is neither history nor
    # "last", and 2024's zero is a block not bearing, not a data point.
    assert ref["last_kg_tree"] is None
    assert ref["avg5_kg_tree"] == 10.0 and ref["seasons_on_file"] == 1
    assert 2025 not in ref["history"]
    # No trees on the register: kg/tree can't be worked out, kg still shows.
    ref = _block_reference({2024: 500.0}, 0, 2025, set(), current_year=2025)
    assert ref["last_kg_tree"] is None and ref["history"][2024]["kg"] == 500.0


def test_estimate_lifecycle(client):
    _seed_history()
    r = client.post("/api/estimate", json={"season_year": 2026, "name": "January"})
    assert r.status_code == 200
    est = r.json()
    # Seeded from the register: own active blocks, current trees, no figure yet.
    assert {l["block_id"]: l["trees"] for l in est["lines"]} == {"7": 2000, "8a": 500}
    assert all(l["kg_per_tree"] is None for l in est["lines"]) and est["total_kg"] == 0

    r = client.put(f"/api/estimate/{est['id']}", json={"lines": [
        {"block_id": "7", "trees": 2000, "kg_per_tree": 64, "note": "can swing a lot"},
        {"block_id": "8a", "trees": 480, "kg_per_tree": 20},
    ]})
    assert r.status_code == 200
    assert r.json()["total_kg"] == 2000 * 64 + 480 * 20

    # A revision copies the last version and becomes the one shown.
    r = client.post("/api/estimate", json={"season_year": 2026, "name": "After fruit set",
                                           "copy_from_id": est["id"]})
    rev = r.json()
    assert rev["total_kg"] == 2000 * 64 + 480 * 20
    client.put(f"/api/estimate/{rev['id']}", json={"lines": [
        {"block_id": "7", "trees": 2000, "kg_per_tree": 40},
        {"block_id": "8a", "trees": 480, "kg_per_tree": 20},
    ]})
    view = client.get("/api/estimate?season=2026").json()
    assert [e["name"] for e in view["estimates"]] == ["After fruit set", "January"]
    assert view["estimate"]["id"] == rev["id"] and view["estimate"]["total_kg"] == 89_600
    # The first version is untouched and still reachable.
    old = client.get(f"/api/estimate?season=2026&estimate_id={est['id']}").json()["estimate"]
    assert old["total_kg"] == 137_600
    note = next(l for l in old["lines"] if l["block_id"] == "7")["note"]
    assert note == "can swing a lot"

    # Estimates are per season.
    assert client.get("/api/estimate?season=2027").json()["estimates"] == []

    r = client.get(f"/api/estimate/{rev['id']}/export")
    assert r.status_code == 200 and r.content[:2] == b"PK"

    assert client.delete(f"/api/estimate/{est['id']}").status_code == 200
    assert client.get(f"/api/estimate?season=2026&estimate_id={est['id']}").status_code == 404
    assert [e["id"] for e in client.get("/api/estimate?season=2026").json()["estimates"]] == [rev["id"]]


def test_estimate_validation(client):
    est = client.post("/api/estimate", json={"season_year": 2026}).json()
    assert est["name"]   # defaults to today's date
    dup = [{"block_id": "7", "trees": 1, "kg_per_tree": 1}, {"block_id": "7", "trees": 1}]
    assert client.put(f"/api/estimate/{est['id']}", json={"lines": dup}).status_code == 400
    neg = [{"block_id": "7", "trees": 10, "kg_per_tree": -1}]
    assert client.put(f"/api/estimate/{est['id']}", json={"lines": neg}).status_code == 422
    assert client.put(f"/api/estimate/{est['id']}", json={"name": "  "}).status_code == 400
    assert client.post("/api/estimate", json={"season_year": 26}).status_code == 400
    assert client.put("/api/estimate/9999", json={"name": "x"}).status_code == 404
    assert client.post("/api/estimate", json={"season_year": 2026, "copy_from_id": 9999}).status_code == 404


def _season(day_kg):
    return {"current_year": 2026, "anchor_month": 8, "anchor_day": 1, "day_kg": day_kg}


def test_progress_projection():
    day_kg = {
        # Two past seasons: 2024 had half its crop in by 1 Oct, 2025 a quarter.
        (2024, "7", date(2024, 9, 20)): 50.0, (2024, "7", date(2024, 11, 20)): 50.0,
        (2025, "7", date(2025, 9, 20)): 25.0, (2025, "7", date(2025, 11, 20)): 75.0,
        (2026, "7", date(2026, 9, 25)): 30_000.0,
    }
    p = _progress(_season(day_kg), 2026, 200_000.0, date(2026, 10, 1))
    assert p["actual_kg"] == 30_000 and p["typical_share"] == 0.375
    assert p["expected_by_now_kg"] == 75_000
    assert p["projected_kg"] == 80_000
    assert p["projected_low_kg"] == 60_000 and p["projected_high_kg"] == 120_000

    # Too early: the typical season was barely picked by this day.
    early = _progress(_season(day_kg), 2026, 200_000.0, date(2026, 9, 1))
    assert early["typical_share"] == 0 and early["projected_kg"] is None

    # Not the current season: no progress at all.
    assert _progress(_season(day_kg), 2027, None, date(2026, 10, 1)) is None


def test_analysis_unchanged_by_shared_helper(client):
    """The per-block history moved into analysis.block_season_kg(); the
    Analysis tab must read the same figures through it."""
    _seed_history()
    body = client.get("/api/analysis/summary").json()
    b7 = next(b for b in body["block_yield"] if b["block_id"] == "7")
    assert b7["by_year"]["2019"] == {"kg": 90000.0, "kg_ha": 22500.0, "kg_tree": 45.0, "annual_only": True}
    assert b7["by_year"]["2025"]["annual_only"] is False
    assert 2019 in body["block_years"] and 2020 in body["historical_years"]


# --------------------------------------------------------------------------- #
# Pack-out (estimation only)
# --------------------------------------------------------------------------- #
def _pack(*rows):
    from types import SimpleNamespace
    return [SimpleNamespace(position=i, channel=c, pack_type=p, kg_per_carton=k, share_pct=s, note="")
            for i, (c, p, k, s) in enumerate(rows)]


def test_packout_arithmetic():
    from routers.estimate import _packout
    po = _packout(400_000, _pack(
        ("Juice", "", None, 14), ("Local", "2 kg", 2.3, 10), ("Local", "PrePac", 2.0, 4),
        ("Export air", "4.5 kg", 4.7, 15), ("Export sea", "4.5 kg", 4.7, 40),
        ("export sea ", "2 kg", 2.2, 12)))
    assert [l["cartons"] for l in po["lines"]] == [None, 17391, 8000, 12766, 34043, 21818]
    assert po["lines"][0]["kg"] == 56_000
    # Channels group case- and space-blind, keeping the first spelling.
    ch = {c["channel"]: c for c in po["channels"]}
    assert set(ch) == {"Juice", "Local", "Export air", "Export sea"}
    assert ch["Export sea"]["kg"] == 208_000 and ch["Export sea"]["cartons"] == 55_861
    assert ch["Local"]["cartons"] == 25_391 and ch["Juice"]["cartons"] is None
    assert po["packed_kg"] == 324_000 and po["not_packed_kg"] == 56_000 and po["cartons"] == 94_018
    assert po["unallocated_pct"] == 5.0 and po["unallocated_kg"] == 20_000
    assert _packout(400_000, []) is None


def test_packout_matches_the_farms_percent_sheet():
    """The OES workbook's '%' sheet: 250 t at its 1996-2000 average shares
    gives 24,546 / 17,586 / 7,841 / 62,271 cartons (row 17)."""
    from routers.estimate import _packout
    po = _packout(250_000, _pack(("Lw", "", 2.3, 22.582), ("Pp Lok", "", 1.8, 12.662),
                                 ("Lug", "", 2.3, 7.2141006), ("See", "", 2.3, 57.2892134)))
    assert [l["cartons"] for l in po["lines"]] == [24546, 17586, 7841, 62271]


_MIX = [
    {"channel": "Juice", "share_pct": 14},
    {"channel": "Export sea", "pack_type": "4.5 kg", "kg_per_carton": 4.7, "share_pct": 40, "note": "HFR"},
    {"channel": "Local", "pack_type": "2 kg", "kg_per_carton": 2.3, "share_pct": 10},
]


def test_packout_lifecycle(client):
    from db import owner_engine as eng
    from models_owner import YieldEstimatePack
    from sqlmodel import select as sel

    est = client.post("/api/estimate", json={"season_year": 2026, "name": "Jan"}).json()
    assert est["pack"] == [] and est["packout"] is None and est["pack_copied_from"] is None
    r = client.put(f"/api/estimate/{est['id']}", json={
        "lines": [{"block_id": "7", "trees": 2000, "kg_per_tree": 50}], "pack": _MIX})
    assert r.status_code == 200
    po = r.json()["packout"]
    assert po["basis_kg"] == 100_000 and po["cartons"] == round(40_000 / 4.7 + 10_000 / 2.3)  # summed, then rounded once
    assert [p["channel"] for p in r.json()["pack"]] == ["Juice", "Export sea", "Local"]
    # Saving the blocks alone keeps the mix, and the pack-out follows the new total.
    r = client.put(f"/api/estimate/{est['id']}", json={
        "lines": [{"block_id": "7", "trees": 2000, "kg_per_tree": 60}]})
    assert len(r.json()["pack"]) == 3 and r.json()["packout"]["basis_kg"] == 120_000
    # A revision copies it...
    rev = client.post("/api/estimate", json={"season_year": 2026, "name": "Rev",
                                             "copy_from_id": est["id"]}).json()
    assert rev["pack"] == r.json()["pack"] and rev["pack_copied_from"]["id"] == est["id"]
    # ...and a fresh estimate for another season starts from the latest mix.
    nxt = client.post("/api/estimate", json={"season_year": 2027}).json()
    assert nxt["pack_copied_from"]["id"] == rev["id"] and len(nxt["pack"]) == 3
    assert nxt["packout"]["basis_kg"] == 0   # no kg/tree yet
    # Deleting a version takes its mix with it.
    client.delete(f"/api/estimate/{rev['id']}")
    with Session(eng) as s:
        assert not s.exec(sel(YieldEstimatePack).where(YieldEstimatePack.estimate_id == rev["id"])).all()
    # The export gains a Pack-out sheet.
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/estimate/{est['id']}/export").content))
    assert wb.sheetnames == ["Estimate", "Pack-out", "Farm by season"]
    total_row = next(r for r in wb["Pack-out"].iter_rows(values_only=True) if r[0] == "Total")
    assert total_row[5] == client.get(f"/api/estimate?season=2026&estimate_id={est['id']}").json()[
        "estimate"]["packout"]["cartons"]


def test_packout_validation(client):
    est = client.post("/api/estimate", json={"season_year": 2026}).json()
    put = lambda pack: client.put(f"/api/estimate/{est['id']}", json={"pack": pack}).status_code
    assert put([{"channel": "A", "share_pct": 60}, {"channel": "B", "share_pct": 40.1}]) == 400
    assert put([{"channel": "A", "share_pct": 60}, {"channel": "B", "share_pct": 40.04}]) == 200
    assert put([{"channel": "Sea", "pack_type": "2 kg", "share_pct": 1},
                {"channel": "sea ", "pack_type": "2 KG", "share_pct": 1}]) == 400
    assert put([{"channel": "Sea", "pack_type": "2 kg", "share_pct": 1},
                {"channel": "Sea", "pack_type": "4.5 kg", "share_pct": 1}]) == 200
    assert put([{"channel": "   ", "share_pct": 1}]) == 400
    assert put([{"channel": "", "share_pct": 1}]) == 422
    assert put([{"channel": "A", "kg_per_carton": 0, "share_pct": 1}]) == 422
    assert put([{"channel": "A", "share_pct": 101}]) == 422
    assert put([{"channel": "A", "share_pct": -1}]) == 422
    assert put([]) == 200
    assert client.get("/api/estimate?season=2026").json()["estimate"]["packout"] is None


# --------------------------------------------------------------------------- #
# Weather-model cross-check
# --------------------------------------------------------------------------- #
def test_crosscheck():
    from routers.estimate import _crosscheck
    assert _crosscheck(420_000, 480_000, 430_000, 350_000) == {
        "gap_kg": -10_000, "gap_pct": -2.3, "position": "within"}
    assert _crosscheck(300_000, 480_000, 430_000, 350_000)["position"] == "below"
    assert _crosscheck(300_000, 480_000, 430_000, 350_000)["gap_pct"] == -30.2
    assert _crosscheck(500_000, 480_000, 430_000, 350_000)["gap_pct"] == 16.3
    assert _crosscheck(500_000, 480_000, 430_000, 350_000)["position"] == "above"
    # Order-blind: a line that slopes the other way swaps the two ends.
    assert _crosscheck(300_000, 350_000, 430_000, 480_000)["position"] == "below"
    assert _crosscheck(420_000, 480_000, 0, 350_000)["gap_pct"] is None
    assert _crosscheck(0, 1, 1, 1) is None and _crosscheck(1, None, 1, 1) is None


_FC = {"season_year": 2026, "built_at": "2026-09-29T06:00:00Z", "favorable_kg": 160_000,
       "expected_kg": 140_000, "unfavorable_kg": 120_000, "live": True, "settled": 1}
_LINES = [{"block_id": "7", "trees": 2000, "kg_per_tree": 64}, {"block_id": "8a", "trees": 480, "kg_per_tree": 20}]


def test_forecast_snapshot_is_saved_with_a_version(client):
    _seed_history()
    est = client.post("/api/estimate", json={"season_year": 2026, "name": "Jan"}).json()
    assert est["forecast_snapshot"] is None
    r = client.put(f"/api/estimate/{est['id']}", json={"lines": _LINES, "forecast": _FC})
    snap = r.json()["forecast_snapshot"]
    assert snap == {"built_at": "2026-09-29T06:00:00Z", "favorable_kg": 160_000, "expected_kg": 140_000,
                    "unfavorable_kg": 120_000, "live": True, "settled": 1,
                    "gap_kg": -2_400, "gap_pct": -1.7, "position": "within"}
    v = client.get("/api/estimate?season=2026").json()
    assert v["estimates"][0]["total_kg"] == 137_600
    assert v["estimates"][0]["forecast"]["expected_kg"] == 140_000
    # A save without one leaves it; a copy doesn't carry it.
    r = client.put(f"/api/estimate/{est['id']}", json={"name": "Jan 2"})
    assert r.json()["forecast_snapshot"]["expected_kg"] == 140_000
    rev = client.post("/api/estimate", json={"season_year": 2026, "copy_from_id": est["id"]}).json()
    assert rev["forecast_snapshot"] is None
    # A zone-less time is taken as UTC.
    r = client.put(f"/api/estimate/{rev['id']}", json={"forecast": {**_FC, "built_at": "2026-09-29T08:00:00"}})
    assert r.json()["forecast_snapshot"]["built_at"] == "2026-09-29T08:00:00Z"
    r = client.put(f"/api/estimate/{rev['id']}", json={"forecast": {**_FC, "built_at": "2026-09-29T10:00:00+02:00"}})
    assert r.json()["forecast_snapshot"]["built_at"] == "2026-09-29T08:00:00Z"


def test_forecast_snapshot_rejections(client):
    est = client.post("/api/estimate", json={"season_year": 2026}).json()
    nxt = client.post("/api/estimate", json={"season_year": 2027}).json()
    put = lambda eid, fc: client.put(f"/api/estimate/{eid}", json={"forecast": fc}).status_code
    assert put(nxt["id"], {**_FC, "season_year": 2027}) == 400   # the model covers the current season only
    assert put(est["id"], {**_FC, "season_year": 2025}) == 400
    assert put(est["id"], {**_FC, "built_at": "2099-01-01T00:00:00Z"}) == 400
    assert put(est["id"], {**_FC, "expected_kg": -1}) == 422
    assert client.post("/api/estimate", json={"season_year": 2027, "forecast": _FC}).status_code == 400
    assert client.post("/api/estimate", json={"season_year": 2026, "forecast": _FC}).json()[
        "forecast_snapshot"]["expected_kg"] == 140_000


def test_season_total(client):
    _seed_history()
    assert client.get("/api/estimate?season=2025").json()["season_total"] == {"kg": 70_000, "partial": False}
    cur = client.get("/api/estimate?season=2026").json()["season_total"]
    assert cur["partial"] is True and cur["kg"] > 0
    assert client.get("/api/estimate?season=2027").json()["season_total"] is None


def test_estimate_never_runs_the_weather_model(client, monkeypatch):
    import routers.risk as risk_module
    boom = lambda *a, **k: (_ for _ in ()).throw(AssertionError("weather model called"))
    monkeypatch.setattr(risk_module, "build_harvest_forecast", boom)
    monkeypatch.setattr(risk_module, "_compute_driver_state", boom)
    monkeypatch.setattr(risk_module, "fetch_forecast_hourly", boom)
    assert client.get("/api/estimate").status_code == 200
    est = client.post("/api/estimate", json={"season_year": 2026}).json()
    assert client.put(f"/api/estimate/{est['id']}", json={"forecast": _FC}).status_code == 200
    body = client.get("/api/estimate").json()
    assert not {"analogs", "factors"} & set(body)


def test_export_keeps_text_that_starts_with_equals_as_text(client):
    import io
    import openpyxl
    est = client.post("/api/estimate", json={"season_year": 2026, "notes": "=hail in Nov"}).json()
    client.put(f"/api/estimate/{est['id']}", json={
        "lines": [{"block_id": "7", "trees": 10, "kg_per_tree": 1, "note": "=same as 8a"}],
        "pack": [{"channel": "=Sea", "pack_type": "=2kg", "share_pct": 50, "note": "=x"}]})
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/estimate/{est['id']}/export").content))
    cells = [(ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if c.data_type == "f"]
    assert cells == []
    texts = {c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row}
    assert {"=hail in Nov", "=same as 8a", "=Sea", "=2kg", "=x"} <= texts


def test_farm_totals_keep_blocks_no_longer_in_the_register(client):
    """A block pulled out since still grew its seasons' fruit: the farm
    total counts it, as the Risk tab does - only the per-block reference
    figures are limited to today's register."""
    with Session(owner_engine) as s:
        s.add(HistoricalAnnualYield(block_id="7", season_year=2018, kg=100_000))
        s.add(HistoricalAnnualYield(block_id="3", season_year=2018, kg=50_000))
        s.add(HistoricalHarvest(block_id="3", harvest_date=date(2022, 11, 1), season_year=2022, kg=7_000))
        s.commit()
    body = client.get("/api/estimate?season=2018").json()
    assert body["season_total"] == {"kg": 150_000, "partial": False}
    farm = {h["year"]: h["kg"] for h in client.get("/api/estimate?season=2026").json()["farm_history"]}
    assert farm[2018] == 150_000 and farm[2022] == 7_000


def test_cartons_round_half_up_like_the_tab():
    from routers.estimate import _packout
    po = _packout(2_000, _pack(("Local", "2 kg", 16, 10)))   # 12.5 cartons
    assert po["lines"][0]["cartons"] == 13 and po["cartons"] == 13 and po["channels"][0]["cartons"] == 13


def test_projection_range_names_its_seasons():
    day_kg = {
        (2024, "7", date(2024, 9, 20)): 50.0, (2024, "7", date(2024, 11, 20)): 50.0,
        (2025, "7", date(2025, 9, 20)): 30.0, (2025, "7", date(2025, 11, 20)): 70.0,
        (2023, "7", date(2023, 11, 20)): 100.0,   # a late start: nothing in by 1 Oct
        (2026, "7", date(2026, 9, 25)): 30_000.0,
    }
    p = _progress(_season(day_kg), 2026, None, date(2026, 10, 1))
    assert p["typical_share"] == round((0.5 + 0.3 + 0) / 3, 3)
    assert p["range_years"] == [2024, 2025] and p["min_range_share"] == 0.1


def test_saving_does_not_scan_the_seasons_crates(client, monkeypatch):
    import routers.estimate as est_module
    boom = lambda *a, **k: (_ for _ in ()).throw(AssertionError("season scan on save"))
    monkeypatch.setattr(est_module, "season_day_kg", boom)
    r = client.post("/api/estimate", json={"season_year": 2026})
    assert r.status_code == 200 and {l["block_id"] for l in r.json()["lines"]} == {"7", "8a"}
    assert client.put(f"/api/estimate/{r.json()['id']}", json={"name": "x", "forecast": _FC}).status_code == 200
