"""CSV/.xlsx parsing for the historical-data imports.

Trimmed from Boord (../Boord/backend/excel_io.py) to just the one function
routers/historical.py needs. Keeping it generic (a header row + data rows in,
list-of-dicts out) means the import endpoints supply no per-file boilerplate.
"""
import csv
import io

from fastapi import UploadFile
from openpyxl import load_workbook


async def parse_uploaded_table(file: UploadFile) -> list[dict]:
    """Returns a list of dict rows keyed by the header row. Accepts .csv or .xlsx."""
    content = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True) if ws is not None else iter(())
        # A blank sheet has no header row to read. next() on it raised
        # StopIteration straight out of this coroutine - a 500 with no word
        # to the owner - where the caller's own "no data rows" 400 is the
        # answer; an empty list gets it there.
        header_row = next(rows_iter, None)
        if header_row is None:
            return []
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        return [dict(zip(headers, row)) for row in rows_iter if any(v is not None for v in row)]
    else:
        # Excel on Windows saves a plain CSV in the PC's ANSI code page
        # (cp1252 here), not UTF-8: a block name or note with an accent then
        # failed to decode and the import answered 500. UTF-8 is tried first
        # because it is strict - cp1252 accepts nearly any byte - so a real
        # UTF-8 file is never misread as Windows text.
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        return [row for row in reader]
