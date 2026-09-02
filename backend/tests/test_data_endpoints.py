"""The data endpoints: auth-gated, wage-free, reading Boord read-only."""
import pytest

from db import get_boord_session
from models_boord import SystemSetting

_RANGE = "period_start=2026-01-01&period_end=2026-12-31"


@pytest.mark.parametrize("path", [
    f"/api/dashboard/summary?{_RANGE}",
    "/api/suppliers",
    "/api/system-settings",
    "/api/lots/pending",
    "/api/lots/in-transit",
    "/api/lots/received",
    "/api/analysis/summary",
    "/api/risk/summary",
])
def test_requires_auth(client, path):
    assert client.get(path).status_code == 401


def test_dashboard_summary_is_wage_free(client, manager_headers):
    r = client.get(f"/api/dashboard/summary?{_RANGE}", headers=manager_headers)
    assert r.status_code == 200
    body = r.json()
    assert set(body) == {"active_teams", "active_workers", "active_blocks", "workers", "blocks"}
    assert "rate_configured" not in body
    assert body["workers"] and all("amount_due" not in w for w in body["workers"])
    # Everything the fake DB dated inside calendar 2026: the own farm's five
    # crates (85 kg), the neighbour's three (300 kg) and the pre-anchor one
    # (50 kg). The Dashboard is the pack house's, so the neighbour counts
    # here - it is the Analysis tab that is the own farm's alone.
    assert body["active_workers"] == 2 and body["active_blocks"] == 3 and body["active_teams"] == 2
    total_kg = sum(w["total_kg"] for w in body["workers"])
    assert round(total_kg, 1) == 435.0


def test_dashboard_supplier_filter(client, manager_headers):
    """The own-farm case is the subtle one: own-farm workers carry
    supplier_id NULL rather than pointing at the "Own Farm" row, so filtering
    by it has to match NULL too or it silently returns nobody."""
    def workers_for(supplier_id):
        r = client.get(f"/api/dashboard/summary?{_RANGE}&supplier_id={supplier_id}",
                       headers=manager_headers)
        assert r.status_code == 200
        return {w["worker_id"]: w["supplier_name"] for w in r.json()["workers"]}

    assert workers_for(1) == {"001": "Own Farm"}       # own farm, worker has NULL supplier_id
    assert workers_for(2) == {"002": "Neighbour Co"}   # an outside supplier
    unfiltered = client.get(f"/api/dashboard/summary?{_RANGE}", headers=manager_headers).json()
    assert {w["worker_id"] for w in unfiltered["workers"]} == {"001", "002"}


def test_dashboard_sorting_and_derived_figures(client, manager_headers):
    body = client.get(f"/api/dashboard/summary?{_RANGE}", headers=manager_headers).json()
    kgs = [w["total_kg"] for w in body["workers"]]
    assert kgs == sorted(kgs, reverse=True), "workers are listed heaviest first"
    names = [b["name"] for b in body["blocks"]]
    assert names == sorted(names, key=str.lower), "blocks are listed by name"
    # 3 crates on block 7, 2 on 8a, each 18 - 1 kg net
    block7 = next(b for b in body["blocks"] if b["block_id"] == "7")
    assert block7["crates"] == 3 and block7["total_kg"] == 51.0
    assert block7["avg_kg_crate"] == 17.0
    assert block7["avg_kg_tree"] == round(51.0 / 2000, 1)
    assert block7["avg_kg_hectare"] == round(51.0 / 4.0, 1)


def test_dashboard_period_is_required(client, manager_headers):
    assert client.get("/api/dashboard/summary", headers=manager_headers).status_code == 422


def test_dashboard_empty_period_is_not_an_error(client, manager_headers):
    body = client.get("/api/dashboard/summary?period_start=2019-01-01&period_end=2019-12-31",
                      headers=manager_headers).json()
    assert body["workers"] == [] and body["blocks"] == []
    assert body["active_teams"] == 0


def test_lots_lists(client, manager_headers):
    assert len(client.get("/api/lots/in-transit", headers=manager_headers).json()) == 1
    received = client.get("/api/lots/received", headers=manager_headers).json()
    assert len(received) == 1 and received[0]["slip_number"] == "261001-003"


def test_system_settings_and_suppliers(client, manager_headers):
    s = client.get("/api/system-settings", headers=manager_headers).json()
    assert s["packhouse_name"] == "Test Farm" and s["current_harvest_year"] == 2026
    # The season anchor has to reach the frontend - the Season preset and the
    # Analysis charts are both drawn from it.
    assert s["season_start_month"] == 8 and s["season_start_day"] == 1
    names = {x["name"] for x in client.get("/api/suppliers", headers=manager_headers).json()}
    assert names == {"Own Farm", "Neighbour Co"}


def test_boord_session_is_read_only():
    gen = get_boord_session()
    session = next(gen)
    try:
        row = session.get(SystemSetting, 1)
        row.packhouse_name = "MUTATED"
        session.add(row)
        with pytest.raises(Exception):  # sqlite3.OperationalError: query_only
            session.commit()
        session.rollback()
    finally:
        gen.close()


def test_analysis_weather_risk_render(client, manager_headers):
    a = client.get("/api/analysis/summary", headers=manager_headers)
    assert a.status_code == 200 and a.json()["current_year"] == 2026
    w = client.get("/api/weather/history", headers=manager_headers)
    assert w.status_code == 200 and "points" in w.json()   # empty owner.db -> empty points, still 200
    rk = client.get("/api/risk/summary", headers=manager_headers)
    assert rk.status_code == 200 and rk.json()["driver_count"] == 4


# --------------------------------------------------------------------------- #
# The season anchor and own-farm scoping (Boord v3.0 onward)
# --------------------------------------------------------------------------- #
def _analysis(client, headers):
    r = client.get("/api/analysis/summary", headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


def test_analysis_reports_the_anchor_it_was_built_against(client, manager_headers):
    """The charts draw a season_day axis and order the monthly heatmap from
    the anchor, so it has to travel with the data - a chart aligned to a
    different anchor than the figures is wrong in a way nothing on screen
    would show."""
    assert _analysis(client, manager_headers)["season_anchor"] == {"month": 8, "day": 1}


def test_season_is_the_anchored_year_not_the_calendar_year(client, manager_headers):
    """The fixture puts one record three weeks BEFORE the 1-August anchor
    (calendar 2026, season 2025) and one AFTER new year (calendar 2027,
    season 2026). A calendar-year filter gets both backwards."""
    data = _analysis(client, manager_headers)
    assert data["current_year"] == 2026

    # 5 own-farm crates at 17 kg net (85.0) + the 40 kg picked in January.
    # The 50 kg picked before the anchor belongs to season 2025 and must not
    # be in here.
    assert data["season_to_date_kg"] == 125.0


def test_analysis_excludes_another_growers_blocks(client, manager_headers):
    """Boord is a pack house: its block register holds the neighbour's
    orchard too. The history this season is compared against is the own
    farm's alone, so the neighbour's 300 kg must not reach any figure here -
    while the Dashboard, which is the pack house's, still shows it."""
    data = _analysis(client, manager_headers)
    block_ids = {b["block_id"] for b in data["block_yield"]}
    assert block_ids == {"7", "8a"}
    # 300 kg of neighbour fruit would be impossible to miss in the total.
    assert data["season_to_date_kg"] == 125.0

    dash = client.get(f"/api/dashboard/summary?{_RANGE}", headers=manager_headers).json()
    assert "90" in {b["block_id"] for b in dash["blocks"]}, \
        "the Dashboard is the pack house's and still shows every supplier"


# --------------------------------------------------------------------------- #
# The Historical Harvest Data workbook
# --------------------------------------------------------------------------- #
def test_historical_harvest_data_workbook(client, manager_headers):
    """The one endpoint that reads both databases at once: Block and
    HarvestRecord from Boord, the two history tables from owner.db. It was
    carried over from Boord as a non-runnable fragment, so this is mostly a
    check that it runs at all - and that it builds the current season fresh
    from live harvest records rather than needing a re-import."""
    import io
    from openpyxl import load_workbook

    r = client.get("/api/reports/historical-harvest-data", headers=manager_headers)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    wb = load_workbook(io.BytesIO(r.content))
    # No history is imported in the fixture, so the only season on file is the
    # current one, built from HarvestRecord. Annual Totals needs
    # HistoricalAnnualYield rows and is correctly absent.
    assert wb.sheetnames == ["Blocks", "Notes", "Season Summary", "Block by Year", "2026"]
    assert "Annual Totals" not in wb.sheetnames

    # The neighbour's block must not appear in a workbook titled as this
    # farm's whole record - same rule as the Analysis tab.
    block_ids = {row[0] for row in wb["Blocks"].iter_rows(min_row=2, values_only=True)}
    assert block_ids == {"7", "8a"}


def test_historical_harvest_data_requires_auth(client):
    assert client.get("/api/reports/historical-harvest-data").status_code == 401
