"""The Markets tab: what the fruit actually made, by pack type, receiver and
day.

Two sources, one per kind of season:

  * a season with a hand-kept "Markte" workbook - MarketPackPrice (season
    totals) and MarketDailySale (a row per delivery), both from
    scripts/import_market_prices.py;
  * the season still being sold, which has no workbook yet - read live from
    Boord Packhouse's delivery notes (what went out, to whom, on which day)
    and Boord's own Markets & Transport records (what the market paid: its
    emailed invoices, or else the statement the office typed). That is the
    same pairing Boord's admin Reports tab makes in
    routers/reports.py:_market_notes, and the records that tab keeps are the
    ones replacing the workbook, so a season shows here as it is sold.

A season that has a workbook takes its daily figures from the workbook alone:
it is the reconciled record, and the live notes for the same days would only
count the same sales twice.

Both live reads are optional, like dashboard.py's Delivered vs Picked:
Packhouse not installed, or a Boord older than its Markets & Transport
tables, just means no live figures. They are deliberately not in
main._BOORD_SCHEMA_CHECK - a farm without them must still boot. Neither
read is held across an outbound call; there is none here.
"""
import io
import json
import os
import re
import sqlite3
from collections import defaultdict
from datetime import date
from typing import List, Optional

import openpyxl
from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from openpyxl.styles import Font
from sqlalchemy import text
from sqlalchemy.exc import OperationalError
from sqlmodel import Session, select

import config
from db import get_boord_session, get_owner_session
from models_boord import SystemSetting
from models_owner import MarketDailyPrice, MarketDailySale, MarketPackPrice
from routers.historical_report import REPORTS_DIR, XLSX_MEDIA, _style_header_cell
from timeutil import season_year_for

router = APIRouter(prefix="/api/market", tags=["market"])


def _packhouse_rows(sql: str, params: tuple = ()) -> list:
    """Rows from Boord Packhouse's packhouse.db as dicts, [] when it isn't
    installed here or hasn't made the table yet."""
    if not os.path.exists(config.PACKHOUSE_DB_PATH):
        return []
    # sqlite3.Error, not just OperationalError: a file that exists but is
    # not a database yet (Packhouse installed, nothing migrated; or mid-copy)
    # raises DatabaseError from connect() or the first statement - and this
    # is read by the Estimate tab as well as the Markets tab, neither of
    # which may fall over for an optional list. See dashboard._delivered_kg.
    try:
        conn = sqlite3.connect(f"file:{config.PACKHOUSE_DB_PATH}?mode=ro", uri=True, timeout=10)
    except sqlite3.Error:
        return []
    conn.row_factory = sqlite3.Row
    try:
        return [dict(r) for r in conn.execute(sql, params).fetchall()]
    except sqlite3.Error:
        return []
    finally:
        conn.close()


def _boord_rows(boord: Session, sql: str) -> list:
    """Rows from one of Boord's Markets & Transport tables, [] on a Boord
    that predates them (no such table)."""
    try:
        return [dict(r) for r in boord.execute(text(sql)).mappings().all()]
    except OperationalError:
        boord.rollback()
        return []


def _pack_key(name: str) -> str:
    """Case/punctuation/whitespace-insensitive key for matching a market
    workbook's own spelling of a pack type against Packhouse's own list -
    the two were never going to agree on hyphens or capitals."""
    return re.sub(r"[^a-z0-9]", "", name.lower())


def _packhouse_names(table: str) -> list:
    """This farm's own active pack types ("packaging", Settings ->
    Consignment) or receivers ("receiver", Settings -> Delivery Note) from
    Boord Packhouse, in Packhouse's own order - the names the Markets tab
    lists, and the spelling a workbook name is shown under when it matches
    one but for case or punctuation. A name Packhouse doesn't have keeps its
    own spelling; never dropped for lack of a match. The ones that differ by
    more than spelling ("Kratte" for "Crate", "Lug" for "Air") are renamed at
    import instead (scripts/import_market_prices.py PACK_RENAMES and
    RECEIVER_RENAMES). Active only: one retired in Packhouse drops out of the
    list, though a season with sales under it still shows them."""
    assert table in ("packaging", "receiver")
    rows = _packhouse_rows(f"SELECT name FROM {table} WHERE active = 1 ORDER BY sort, lower(name)")
    return [r["name"] for r in rows]


def _canonical(names: list) -> dict:
    return {_pack_key(n): n for n in names}


def _season_totals(owner: Session, canonical: dict) -> dict:
    """Each pack type's season-wide R/carton from MarketPackPrice, weighted
    across buyers - sum(sub_total)/sum(cartons_sold), so a buyer with ten
    cartons doesn't move it as much as one with ten thousand. Two workbook
    spellings that align to one Packhouse name are summed as one."""
    totals: dict = defaultdict(lambda: defaultdict(lambda: [0.0, 0.0]))  # season -> pack -> [sub_total, cartons]
    for r in owner.exec(select(MarketPackPrice)).all():
        acc = totals[r.season_year][canonical.get(_pack_key(r.pack_type), r.pack_type)]
        acc[0] += r.sub_total
        acc[1] += r.cartons_sold
    return {
        season: {pack: {"cartons": round(cartons, 1), "rand_per_carton": round(sub / cartons, 2)}
                 for pack, (sub, cartons) in packs.items() if cartons}
        for season, packs in totals.items()
    }


def _live_sales(boord: Session) -> list:
    """(sale_date, receiver, pack_type, cartons_sold, gross) per pack per
    delivery note, from Packhouse's notes and Boord's market records - the
    money from the market's invoices once any have arrived for the note (all
    its part-payments added up), else the office's typed statement. A note
    with neither is not sold yet and adds nothing. Juice notes (bins, not
    pallets) are not a pack type and are left out."""
    notes = _packhouse_rows("SELECT uuid, note_no, loaded_on, receiver, bins FROM deliverynote "
                            "WHERE removed = 0 AND loaded_on != ''")
    if not notes:
        return []
    statements = {r["note_uuid"]: json.loads(r["lines"] or "[]")
                  for r in _boord_rows(boord, "SELECT note_uuid, lines FROM marketstatement")}
    invoices = defaultdict(list)
    for r in _boord_rows(boord, "SELECT note_no, lines FROM marketinvoice"):
        invoices[r["note_no"]] += json.loads(r["lines"] or "[]")
    pack_of = {r["product"]: r["pack"] for r in _boord_rows(boord, "SELECT product, pack FROM marketproduct")}

    out = []
    for n in notes:
        if json.loads(n["bins"] or "[]"):
            continue
        try:
            sale_date = date.fromisoformat(n["loaded_on"])
        except ValueError:
            continue
        packs: dict = defaultdict(lambda: [0.0, 0.0])   # pack -> [cartons, gross]
        if invoices.get(str(n["note_no"])):
            for ln in invoices[str(n["note_no"])]:
                acc = packs[pack_of.get(ln["product"], ln["product"])]
                acc[0] += ln.get("sold") or 0
                acc[1] += ln.get("income") or 0
        else:
            for ln in statements.get(n["uuid"], []):
                if ln.get("sold") and ln.get("price"):
                    acc = packs[ln["pack"]]
                    acc[0] += ln["sold"]
                    acc[1] += ln["sold"] * ln["price"]
        receiver = (n["receiver"] or "").split("\n")[0].strip()
        out += [(sale_date, receiver, pack, cartons, gross) for pack, (cartons, gross) in packs.items() if cartons]
    return out


def _daily(owner: Session, boord: Session, packs: list, receivers: list) -> dict:
    """The daily chart's points: one per season, day, receiver and pack
    type, R/carton weighted by cartons when a day had several lines (two
    cultivars on one note, two notes the same day) - and Joburg Market's
    daily figures beside them, under `reference`. A season is offered when
    either has a day in it.

    `packs` and `receivers` are what the tab lists: Packhouse's active ones
    first, in Packhouse's order, so a pack type or receiver set up there
    during the season is offered at once and charts as soon as it sells -
    then any other name with sales on file (a receiver since retired).
    `season_anchor` is Boord's season start, which the chart lines seasons
    up by: day N of one season sits over day N of every other."""
    pack_of, receiver_of = _canonical(packs), _canonical(receivers)
    acc: dict = defaultdict(lambda: [0.0, 0.0])   # (season, day, receiver, pack) -> [cartons, gross]
    workbook_seasons = set()
    for r in owner.exec(select(MarketDailySale)).all():
        workbook_seasons.add(r.season_year)
        a = acc[(r.season_year, r.sale_date, receiver_of.get(_pack_key(r.receiver), r.receiver),
                 pack_of.get(_pack_key(r.pack_type), r.pack_type))]
        a[0] += r.cartons_sold
        a[1] += r.cartons_sold * r.price_per_carton

    settings = boord.exec(select(SystemSetting)).first()
    month, day = (settings.season_start_month, settings.season_start_day) if settings else (1, 1)
    for sale_date, receiver, pack, cartons, gross in _live_sales(boord):
        season = season_year_for(sale_date, month, day)
        if season in workbook_seasons:
            continue
        a = acc[(season, sale_date, receiver_of.get(_pack_key(receiver), receiver),
                 pack_of.get(_pack_key(pack), pack))]
        a[0] += cartons
        a[1] += gross

    points = [{"season_year": s, "date": d.isoformat(), "receiver": rcv, "pack_type": pk,
               "cartons": round(c, 1), "price": round(g / c, 2)}
              for (s, d, rcv, pk), (c, g) in sorted(acc.items()) if c]
    # Joburg Market's whole-market Litchis total, the chart's dashed line: the
    # average price per unit sold (value over units - its litchi units are
    # mostly 2kg boxes, so it reads beside R/carton) and the units sold. A day
    # the page gave no unit count has no unit price, but still its kg figures.
    reference = [{"season_year": season_year_for(r.price_date, month, day), "date": r.price_date.isoformat(),
                  "price": round(r.value_sold / r.qty_sold, 2) if r.qty_sold else None,
                  "units": r.qty_sold, "rand_per_kg": r.rand_per_kg, "kg": r.kg_sold}
                 for r in owner.exec(select(MarketDailyPrice).order_by(MarketDailyPrice.price_date)).all()]
    return {"seasons": sorted({p["season_year"] for p in points + reference}),
            "workbook_seasons": sorted(workbook_seasons),
            "season_anchor": {"month": month, "day": day},
            "packs": packs + sorted({p["pack_type"] for p in points} - set(packs)),
            "receivers": receivers + sorted({p["receiver"] for p in points} - set(receivers)),
            "points": points, "reference": reference}


@router.get("/summary")
def market_summary(boord: Session = Depends(get_boord_session),
                   owner: Session = Depends(get_owner_session)):
    """Everything the Markets tab shows: each pack type's season R/carton
    (the table), the daily points per receiver and pack type (the chart),
    and the latest Joburg Market reference price for Litchis as a whole
    (market_reference.py) - a free, whole-market figure, never broken down
    by pack type the way the farm's own figures are."""
    packs, receivers = _packhouse_names("packaging"), _packhouse_names("receiver")
    by_season = _season_totals(owner, _canonical(packs))
    ref = owner.exec(select(MarketDailyPrice).order_by(MarketDailyPrice.price_date.desc())).first()
    return {
        "seasons": sorted(by_season),
        "pack_types": sorted({pt for packs in by_season.values() for pt in packs}),
        "by_season": by_season,
        "daily": _daily(owner, boord, packs, receivers),
        "joburg_market_reference": ({"date": ref.price_date.isoformat(), "rand_per_kg": ref.rand_per_kg,
                                     "kg_sold": ref.kg_sold} if ref else None),
    }


def _header(ws, cells: list) -> None:
    ws.append(cells)
    for c in ws[ws.max_row]:
        _style_header_cell(c)


@router.get("/export")
def market_export(season: Optional[List[int]] = Query(None),
                  pack: Optional[List[str]] = Query(None), receiver: Optional[List[str]] = Query(None),
                  boord: Session = Depends(get_boord_session), owner: Session = Depends(get_owner_session)):
    """The Markets tab as a workbook. "Daily" is what the chart shows: the
    seasons ticked (`season`, repeated; else the latest with daily figures),
    cut to the pack type and receivers ticked (`pack`/`receiver`, repeated;
    absent = all) - a row per season, day, receiver and pack type with week,
    cartons, R/carton and gross, filterable in Excel. "Season totals" is the
    table below the chart, and "Joburg Market" every daily reference price
    this app has kept, since the market's own site keeps none."""
    packs, receivers = _packhouse_names("packaging"), _packhouse_names("receiver")
    daily = _daily(owner, boord, packs, receivers)
    seasons = sorted(set(season or []) & set(daily["seasons"])) or daily["seasons"][-1:]
    points = [p for p in daily["points"] if p["season_year"] in seasons
              and (pack is None or p["pack_type"] in pack) and (receiver is None or p["receiver"] in receiver)]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily"
    names = " and ".join(str(s) for s in seasons)
    ws.append([f"Daily market sales, {names} season{'s' if len(seasons) > 1 else ''}" if seasons
               else "Daily market sales - none on file yet"])
    ws["A1"].font = Font(bold=True, size=14)
    sources = "; ".join(
        f"{s} from the farm's market workbook" if s in daily["workbook_seasons"]
        else f"{s} from Boord Packhouse's delivery notes, priced from Boord's market invoices and statements"
        for s in seasons)
    ws.append([f"{sources or 'No daily figures on file'}. Pack type: {', '.join(pack) if pack is not None else 'all'}. "
               f"Receivers: {', '.join(receiver) if receiver is not None else 'all'}."])
    ws.append([])
    _header(ws, ["Season", "Date", "Week", "Receiver", "Pack type", "Cartons sold", "R/carton", "Gross (R)"])
    header_row = ws.max_row
    for p in points:
        d = date.fromisoformat(p["date"])
        ws.append([p["season_year"], d, d.isocalendar()[1], p["receiver"], p["pack_type"], p["cartons"],
                   p["price"], round(p["cartons"] * p["price"], 2)])
    for r in range(header_row + 1, ws.max_row + 1):
        ws.cell(row=r, column=2).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=7).number_format = "#,##0.00"
        ws.cell(row=r, column=8).number_format = "#,##0.00"
    ws.auto_filter.ref = f"A{header_row}:H{ws.max_row}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for col, width in zip("ABCDEFGH", (8, 12, 7, 22, 16, 13, 11, 14)):
        ws.column_dimensions[col].width = width

    by_season = _season_totals(owner, _canonical(packs))
    total_seasons = sorted(by_season, reverse=True)
    ts = wb.create_sheet("Season totals")
    ts.append(["R per carton by pack type, weighted by cartons sold across every buyer"])
    ts["A1"].font = Font(bold=True)
    ts.append([])
    _header(ts, ["Pack type"] + [f"{s} R/carton" for s in total_seasons] + [f"{s} cartons" for s in total_seasons])
    for pt in sorted({pt for season_packs in by_season.values() for pt in season_packs}):
        lines = [by_season[s].get(pt) for s in total_seasons]
        ts.append([pt] + [ln["rand_per_carton"] if ln else None for ln in lines]
                  + [ln["cartons"] if ln else None for ln in lines])
    ts.column_dimensions["A"].width = 16

    if daily["reference"]:
        js = wb.create_sheet("Joburg Market")
        js.append(["Joburg Market's own daily Litchis total - whole market, every pack and agent, recorded "
                   "by this app each trading day (the market keeps no history of its own)"])
        js["A1"].font = Font(bold=True)
        js.append([])
        _header(js, ["Date", "Week", "Season", "Avg R/unit", "Units sold", "R/kg", "Kg sold"])
        for r in daily["reference"]:
            d = date.fromisoformat(r["date"])
            js.append([d, d.isocalendar()[1], r["season_year"], r["price"], r["units"], r["rand_per_kg"], r["kg"]])
            js.cell(row=js.max_row, column=1).number_format = "dd/mm/yyyy"
            js.cell(row=js.max_row, column=4).number_format = "#,##0.00"
            js.cell(row=js.max_row, column=6).number_format = "#,##0.00"
        js.column_dimensions["A"].width = 12

    # Receiver names come from Boord's delivery notes, typed by people: one
    # starting with "=" would open as a formula. Nothing here is meant to be.
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    filename = f"Market_Prices_{'-'.join(str(s) for s in seasons)}.xlsx" if seasons else "Market_Prices.xlsx"
    # Kept on disk as well, like every workbook this app builds - see
    # routers/historical_report.py's REPORTS_DIR.
    with open(os.path.join(REPORTS_DIR, filename), "wb") as f:
        f.write(data)
    return Response(data, media_type=XLSX_MEDIA, headers={"Content-Disposition": f'attachment; filename="{filename}"'})
