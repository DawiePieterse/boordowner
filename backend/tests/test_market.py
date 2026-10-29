"""The Markets tab (/api/market/summary, routers/market.py): MarketPackPrice
and MarketDailySale from the farm's hand-kept litchi market workbooks
(scripts/import_market_prices.py), the season still being sold read live
from Packhouse's delivery notes and Boord's market records, and the Joburg
Market daily reference price (market_reference.py). Tests never touch the
network: market_reference's own HTML parsing is exercised against canned
strings, not a live fetch."""
import json
import os
import sqlite3
from datetime import date, datetime

from sqlmodel import Session, select

import config
from db import owner_engine
from models_owner import MarketDailyPrice, MarketDailySale, MarketPackPrice


def _market(client):
    r = client.get("/api/market/summary")
    assert r.status_code == 200, r.text
    return r.json()


def _write_packaging(names):
    """names: [(id, name, sort, active), ...] into a fresh packhouse.db's
    packaging table at config.PACKHOUSE_DB_PATH - same shape as
    test_data_endpoints.py's _write_packhouse_db. Each test that writes one
    cleans it up itself."""
    conn = sqlite3.connect(config.PACKHOUSE_DB_PATH)
    conn.execute("CREATE TABLE packaging (id INTEGER, name TEXT, sort INTEGER, active INTEGER)")
    conn.executemany("INSERT INTO packaging (id, name, sort, active) VALUES (?, ?, ?, ?)", names)
    conn.commit()
    conn.close()


def test_market_summary_empty_when_nothing_imported(client):
    """Nothing imported, no packhouse.db, and the fake Boord has no Markets &
    Transport tables (an older Boord) - every part empty, never a 500."""
    market = _market(client)
    assert market == {"seasons": [], "pack_types": [], "by_season": {},
                      "daily": {"seasons": [], "workbook_seasons": [], "season_anchor": {"month": 8, "day": 1},
                                "packs": [], "receivers": [], "points": [], "reference": []},
                      "joburg_market_reference": None}


def test_market_summary_weights_by_cartons_across_buyers(client):
    """Two buyers sold the same pack type in the same season at different
    prices and volumes - the season figure must be sub_total-weighted
    (sum/sum), not a plain average of the two buyers' prices, so a buyer
    with ten cartons doesn't move it as much as one with ten thousand."""
    with Session(owner_engine) as owner:
        owner.add_all([
            MarketPackPrice(season_year=2025, buyer="Dapper", pack_type="2kg XL",
                            cartons_sold=100, price_per_carton=200.0, sub_total=20000.0),
            MarketPackPrice(season_year=2025, buyer="ToroFruit", pack_type="2kg XL",
                            cartons_sold=900, price_per_carton=300.0, sub_total=270000.0),
            MarketPackPrice(season_year=2024, buyer="Dapper", pack_type="2kg XL",
                            cartons_sold=50, price_per_carton=100.0, sub_total=5000.0),
        ])
        owner.commit()
    market = _market(client)
    assert market["seasons"] == [2024, 2025]
    assert market["pack_types"] == ["2kg XL"]
    # (20000 + 270000) / (100 + 900) = 290.0
    assert market["by_season"]["2025"]["2kg XL"] == {"cartons": 1000.0, "rand_per_carton": 290.0}
    assert market["by_season"]["2024"]["2kg XL"] == {"cartons": 50.0, "rand_per_carton": 100.0}


def test_market_summary_skips_a_pack_type_with_no_cartons(client):
    """A row somehow imported with 0 cartons (shouldn't happen - the import
    script already filters these - but division by zero must never 500 the
    Analysis tab over one bad row)."""
    with Session(owner_engine) as owner:
        owner.add(MarketPackPrice(season_year=2025, buyer="Dapper", pack_type="PrePac",
                                  cartons_sold=0, price_per_carton=0.0, sub_total=0.0))
        owner.commit()
    market = _market(client)
    assert market["by_season"] == {"2025": {}}


# --------------------------------------------------------------------------- #
# Aligning pack-type spelling with Boord Packhouse's own list
# (routers/analysis.py _packhouse_pack_names / _pack_key).
# --------------------------------------------------------------------------- #
def test_market_summary_aligns_to_packhouses_own_spelling(client):
    """The workbook calls it "PrePac"; Packhouse's own Settings ->
    Consignment list (the farm's one canonical naming) spells it
    "Pre-pac" - the figure must come out under Packhouse's spelling."""
    _write_packaging([(1, "Pre-pac", 0, 1)])
    try:
        with Session(owner_engine) as owner:
            owner.add(MarketPackPrice(season_year=2025, buyer="Dapper", pack_type="PrePac",
                                      cartons_sold=100, price_per_carton=200.0, sub_total=20000.0))
            owner.commit()
        market = _market(client)
        assert market["pack_types"] == ["Pre-pac"]
        assert market["by_season"]["2025"] == {"Pre-pac": {"cartons": 100.0, "rand_per_carton": 200.0}}
    finally:
        os.remove(config.PACKHOUSE_DB_PATH)


def test_market_summary_merges_rows_that_align_to_the_same_packhouse_name(client):
    """Two differently-spelled workbook rows ("PrePac", "Pre Pac") that both
    align to the one Packhouse pack type must be weighted and summed
    together, not shown as two separate columns."""
    _write_packaging([(1, "Pre-pac", 0, 1)])
    try:
        with Session(owner_engine) as owner:
            owner.add_all([
                MarketPackPrice(season_year=2025, buyer="Dapper", pack_type="PrePac",
                                cartons_sold=100, price_per_carton=200.0, sub_total=20000.0),
                MarketPackPrice(season_year=2025, buyer="ToroFruit", pack_type="Pre Pac",
                                cartons_sold=50, price_per_carton=100.0, sub_total=5000.0),
            ])
            owner.commit()
        market = _market(client)
        assert market["pack_types"] == ["Pre-pac"]
        # (20000 + 5000) / (100 + 50) = 166.67
        assert market["by_season"]["2025"]["Pre-pac"] == {"cartons": 150.0, "rand_per_carton": 166.67}
    finally:
        os.remove(config.PACKHOUSE_DB_PATH)


def test_market_summary_keeps_workbook_spelling_without_a_packhouse_match(client):
    """A pack type Packhouse hasn't set up yet must never be dropped - it
    keeps the workbook's own spelling."""
    _write_packaging([(1, "Pre-pac", 0, 1)])
    try:
        with Session(owner_engine) as owner:
            owner.add(MarketPackPrice(season_year=2025, buyer="Dapper", pack_type="2kg XL",
                                      cartons_sold=100, price_per_carton=200.0, sub_total=20000.0))
            owner.commit()
        market = _market(client)
        assert market["pack_types"] == ["2kg XL"]
    finally:
        os.remove(config.PACKHOUSE_DB_PATH)


def test_market_summary_ignores_an_inactive_packhouse_pack_type(client):
    """A pack type the farm retired in Packhouse's Settings (active=0,
    never deleted - see models.Packaging) must not pull a current workbook
    row under a name nobody offers any more."""
    _write_packaging([(1, "Pre-pac", 0, 0)])
    try:
        with Session(owner_engine) as owner:
            owner.add(MarketPackPrice(season_year=2025, buyer="Dapper", pack_type="PrePac",
                                      cartons_sold=100, price_per_carton=200.0, sub_total=20000.0))
            owner.commit()
        market = _market(client)
        assert market["pack_types"] == ["PrePac"]
    finally:
        os.remove(config.PACKHOUSE_DB_PATH)


def test_market_summary_without_packhouse_db_keeps_workbook_spelling(client):
    """The common case (Packhouse not installed on this server, or
    unreachable): no alignment happens, same as dashboard.py's
    delivered_kg - the rest of the figure must still come through."""
    assert not os.path.exists(config.PACKHOUSE_DB_PATH)
    with Session(owner_engine) as owner:
        owner.add(MarketPackPrice(season_year=2025, buyer="Dapper", pack_type="PrePac",
                                  cartons_sold=100, price_per_carton=200.0, sub_total=20000.0))
        owner.commit()
    market = _market(client)
    assert market["pack_types"] == ["PrePac"]


# --------------------------------------------------------------------------- #
# The daily chart's points: the workbook's per-receiver rows, and the season
# still being sold read live (routers/market.py _daily / _live_sales).
# --------------------------------------------------------------------------- #
def test_daily_points_weight_a_days_lines_by_cartons(client):
    """Two lines the same day, receiver and pack (two cultivars on one note)
    are one point, its price weighted by cartons: (100x200 + 300x300)/400."""
    with Session(owner_engine) as owner:
        owner.add_all([
            MarketDailySale(season_year=2025, sale_date=date(2025, 10, 29), receiver="ToroFruit",
                            note_no="1259", pack_type="2kg XL", cartons_sold=100, price_per_carton=200.0),
            MarketDailySale(season_year=2025, sale_date=date(2025, 10, 29), receiver="ToroFruit",
                            note_no="1259", pack_type="2kg XL", cartons_sold=300, price_per_carton=300.0),
        ])
        owner.commit()
    daily = _market(client)["daily"]
    assert daily["points"] == [{"season_year": 2025, "date": "2025-10-29", "receiver": "ToroFruit",
                                "pack_type": "2kg XL", "cartons": 400.0, "price": 275.0}]
    assert daily["workbook_seasons"] == [2025]


def _write_live_season(notes, statements=(), invoices=(), products=()):
    """Packhouse delivery notes into packhouse.db, and Boord's Markets &
    Transport tables into the fake boord.db - which, like an older Boord,
    doesn't have them. Undone by _drop_live_season."""
    conn = sqlite3.connect(config.PACKHOUSE_DB_PATH)
    conn.execute("CREATE TABLE deliverynote (uuid TEXT, note_no INTEGER, loaded_on TEXT, receiver TEXT, "
                 "bins TEXT, removed INTEGER)")
    conn.executemany("INSERT INTO deliverynote VALUES (?, ?, ?, ?, ?, 0)", notes)
    conn.commit()
    conn.close()
    conn = sqlite3.connect(config.BOORD_DB_PATH)
    conn.execute("CREATE TABLE marketstatement (note_uuid TEXT, lines TEXT)")
    conn.execute("CREATE TABLE marketinvoice (note_no TEXT, lines TEXT)")
    conn.execute("CREATE TABLE marketproduct (product TEXT, pack TEXT)")
    conn.executemany("INSERT INTO marketstatement VALUES (?, ?)", statements)
    conn.executemany("INSERT INTO marketinvoice VALUES (?, ?)", invoices)
    conn.executemany("INSERT INTO marketproduct VALUES (?, ?)", products)
    conn.commit()
    conn.close()


def _drop_live_season():
    os.remove(config.PACKHOUSE_DB_PATH)
    conn = sqlite3.connect(config.BOORD_DB_PATH)
    for t in ("marketstatement", "marketinvoice", "marketproduct"):
        conn.execute(f"DROP TABLE IF EXISTS {t}")
    conn.commit()
    conn.close()


_LIVE_NOTES = [
    ("n1", 1301, "2026-11-03", "Toro Fruit\nDurban Market", "[]"),
    ("n2", 1302, "2026-11-03", "Dapper", "[]"),
    ("j1", 1, "2026-11-03", "Onderberg", '[{"container": "bin"}]'),   # juice: not a pack type
    ("n3", 1303, "2026-11-04", "Dapper", "[]"),                      # loaded, not sold yet
]


def test_live_season_comes_from_delivery_notes_and_market_records(client):
    """A note's money comes from the market's invoices once there are any -
    the typed statement for n2 (999/carton) loses to its invoice - and else
    from the statement; a pack with no sold/price typed adds nothing, nor do
    the juice note and the unsold one. The receiver is the note's first
    line. 2026-11 is season 2026 under the fixture's 1-August anchor."""
    _write_live_season(
        _LIVE_NOTES,
        statements=[("n1", json.dumps([{"pack": "2kg XL", "sold": 100, "price": 250},
                                       {"pack": "PrePac", "sold": None, "price": None}])),
                    ("n2", json.dumps([{"pack": "2kg XL", "sold": 10, "price": 999}]))],
        invoices=[("1302", json.dumps([{"product": "LICH, BX020, XL", "sold": 20, "income": 5000}]))],
        products=[("LICH, BX020, XL", "2kg XL")])
    try:
        daily = _market(client)["daily"]
    finally:
        _drop_live_season()
    assert daily["points"] == [
        {"season_year": 2026, "date": "2026-11-03", "receiver": "Dapper", "pack_type": "2kg XL",
         "cartons": 20.0, "price": 250.0},
        {"season_year": 2026, "date": "2026-11-03", "receiver": "Toro Fruit", "pack_type": "2kg XL",
         "cartons": 100.0, "price": 250.0},
    ]
    assert daily["seasons"] == [2026] and daily["workbook_seasons"] == []


def test_joburg_days_ride_with_the_daily_points_and_open_their_own_season(client):
    """Price is per unit sold (value over units), filed under its season by
    Boord's anchor (1 August in the fixture) - and a season with only Joburg
    days is still offered, so the line shows before the farm's first sale."""
    with Session(owner_engine) as owner:
        owner.add_all([
            MarketDailyPrice(price_date=date(2026, 10, 28), value_sold=90200.0, qty_sold=312.0, kg_sold=658.0,
                             rand_per_kg=137.08, fetched_at=datetime(2026, 10, 28, 13)),
            MarketDailyPrice(price_date=date(2026, 10, 29), value_sold=5000.0, qty_sold=None, kg_sold=40.0,
                             rand_per_kg=125.0, fetched_at=datetime(2026, 10, 29, 13)),
        ])
        owner.commit()
    daily = _market(client)["daily"]
    assert daily["seasons"] == [2026] and daily["points"] == []
    assert daily["reference"] == [
        {"season_year": 2026, "date": "2026-10-28", "price": 289.1, "units": 312.0, "rand_per_kg": 137.08, "kg": 658.0},
        {"season_year": 2026, "date": "2026-10-29", "price": None, "units": None, "rand_per_kg": 125.0, "kg": 40.0},
    ]


def test_a_season_with_a_workbook_ignores_its_live_notes(client):
    """The workbook is the reconciled record: once a season has one, its
    live notes would only count the same sales twice."""
    with Session(owner_engine) as owner:
        owner.add(MarketDailySale(season_year=2026, sale_date=date(2026, 11, 3), receiver="ToroFruit",
                                  pack_type="2kg XL", cartons_sold=50, price_per_carton=240.0))
        owner.commit()
    _write_live_season(_LIVE_NOTES, statements=[("n1", json.dumps([{"pack": "2kg XL", "sold": 100, "price": 250}]))])
    try:
        points = _market(client)["daily"]["points"]
    finally:
        _drop_live_season()
    assert [(p["receiver"], p["cartons"]) for p in points] == [("ToroFruit", 50.0)]


# --------------------------------------------------------------------------- #
# The Excel export (routers/market.py market_export)
# --------------------------------------------------------------------------- #
def _export(client, query=""):
    import io

    import openpyxl
    r = client.get(f"/api/market/export{query}")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return r.headers["content-disposition"], openpyxl.load_workbook(io.BytesIO(r.content))


def _daily_rows(ws):
    """The Daily sheet's data rows, below its header (row 4)."""
    return [row for row in ws.iter_rows(min_row=5, values_only=True)]


def _add_sales(*rows):
    with Session(owner_engine) as owner:
        owner.add_all([MarketDailySale(season_year=s, sale_date=d, receiver=rcv, pack_type=pk,
                                       cartons_sold=c, price_per_carton=p) for s, d, rcv, pk, c, p in rows])
        owner.commit()


def test_export_cuts_the_daily_sheet_to_what_is_ticked(client):
    _add_sales((2025, date(2025, 10, 29), "ToroFruit", "2kg XL", 100, 220.0),
               (2025, date(2025, 10, 29), "Dapper", "2kg XL", 10, 350.0),
               (2025, date(2025, 10, 30), "ToroFruit", "PrePac", 50, 230.0))
    disposition, wb = _export(client, "?season=2025&receiver=ToroFruit&pack=2kg%20XL")
    assert 'filename="Market_Prices_2025.xlsx"' in disposition
    assert _daily_rows(wb["Daily"]) == [(2025, datetime(2025, 10, 29), 44, "ToroFruit", "2kg XL", 100, 220, 22000)]
    assert "Receivers: ToroFruit" in wb["Daily"]["A2"].value


def test_export_defaults_to_the_latest_season_and_everything(client):
    _add_sales((2024, date(2024, 11, 5), "Dapper", "PrePac", 5, 500.0),
               (2025, date(2025, 10, 29), "ToroFruit", "2kg XL", 100, 220.0),
               (2025, date(2025, 10, 29), "Dapper", "2kg XL", 10, 350.0))
    disposition, wb = _export(client)
    assert 'filename="Market_Prices_2025.xlsx"' in disposition
    assert [r[3] for r in _daily_rows(wb["Daily"])] == ["Dapper", "ToroFruit"]
    assert "Joburg Market" not in wb.sheetnames   # none kept yet - no empty sheet


def test_export_takes_several_seasons(client):
    """The chart overlays the seasons ticked, so the workbook carries them all."""
    _add_sales((2024, date(2024, 11, 5), "Dapper", "PrePac", 5, 500.0),
               (2025, date(2025, 10, 29), "Dapper", "PrePac", 10, 350.0))
    disposition, wb = _export(client, "?season=2024&season=2025")
    assert 'filename="Market_Prices_2024-2025.xlsx"' in disposition
    assert [(r[0], r[3]) for r in _daily_rows(wb["Daily"])] == [(2024, "Dapper"), (2025, "Dapper")]
    assert wb["Daily"]["A1"].value == "Daily market sales, 2024 and 2025 seasons"


def _write_receivers(names):
    """names: [(name, sort, active), ...] into packhouse.db's receiver
    table (Settings -> Delivery Note), beside any packaging table already
    written. The test removes packhouse.db itself."""
    conn = sqlite3.connect(config.PACKHOUSE_DB_PATH)
    conn.execute("CREATE TABLE receiver (name TEXT, sort INTEGER, active INTEGER)")
    conn.executemany("INSERT INTO receiver VALUES (?, ?, ?)", names)
    conn.commit()
    conn.close()


def test_lists_follow_packhouse_and_receivers_take_its_spelling(client):
    """The tab lists Packhouse's active pack types and receivers in its own
    order - one set up during the season is offered before it has sold
    anything - then any other name with sales on file. A workbook receiver
    differing only by case or punctuation ("AGRILINK usa-air") is shown
    under Packhouse's spelling; a retired one (active=0) is not listed."""
    _write_packaging([(1, "PrePac", 0, 1), (2, "2kg XL", 1, 1)])
    _write_receivers([("Agrilink USA Air", 0, 1), ("Tomahawk", 1, 1), ("Old Agent", 2, 0)])
    try:
        _add_sales((2025, date(2025, 10, 29), "AGRILINK usa-air", "2kg XL", 10, 300.0),
                   (2025, date(2025, 10, 30), "Dapper", "15kg Crate", 5, 500.0))
        daily = _market(client)["daily"]
    finally:
        os.remove(config.PACKHOUSE_DB_PATH)
    assert daily["packs"] == ["PrePac", "2kg XL", "15kg Crate"]
    assert daily["receivers"] == ["Agrilink USA Air", "Tomahawk", "Dapper"]
    assert [p["receiver"] for p in daily["points"]] == ["Agrilink USA Air", "Dapper"]


def test_export_keeps_a_typed_formula_as_text(client):
    """A receiver typed as "=1+1" must open as that text, not run."""
    _add_sales((2025, date(2025, 10, 29), "=1+1", "2kg XL", 10, 300.0))
    _, wb = _export(client)
    cell = wb["Daily"].cell(row=5, column=4)
    assert cell.value == "=1+1" and cell.data_type == "s"


def test_export_with_nothing_on_file_still_opens(client):
    disposition, wb = _export(client)
    assert 'filename="Market_Prices.xlsx"' in disposition
    assert _daily_rows(wb["Daily"]) == [] and "Season totals" in wb.sheetnames


def test_market_reference_surfaces_the_latest_day(client):
    with Session(owner_engine) as owner:
        owner.add_all([
            MarketDailyPrice(price_date=date(2026, 10, 27), value_sold=80000.0, qty_sold=300.0,
                             kg_sold=600.0, rand_per_kg=133.33, fetched_at=datetime(2026, 10, 27, 13, 0, 0)),
            MarketDailyPrice(price_date=date(2026, 10, 28), value_sold=90200.0, qty_sold=312.0,
                             kg_sold=658.0, rand_per_kg=137.08, fetched_at=datetime(2026, 10, 28, 13, 0, 0)),
        ])
        owner.commit()
    reference = _market(client)["joburg_market_reference"]
    assert reference == {"date": "2026-10-28", "rand_per_kg": 137.08, "kg_sold": 658.0}


# --------------------------------------------------------------------------- #
# market_reference.py's own HTML parsing - no network in tests, so this is
# exercised against canned snippets shaped like the real page's table.
# --------------------------------------------------------------------------- #
import market_reference  # noqa: E402

_TABLE_HTML = """
<p>Please note that prices on our website are updated between 12 noon and 1pm every week day.</p>
<p>This information is for <b>24 October 2025</b> </p>
<table>
<tr><th>Commodity</th><th>Total Value Sold</th><th>Total Qty Sold</th><th>Total Kg Sold</th><th>Qty Available</th></tr>
<tr><td>APPLES</td><td>R2,516,128.43<br>MTD: R56,319,596.70</td><td>17,397<br>MTD: 394,759</td>
    <td>215,688<br>MTD: 5,009,277</td><td>63,047</td></tr>
<tr><td>LITCHIS</td><td>R90,200.00<br>MTD: R525,810.00</td><td>312<br>MTD: 1,910</td>
    <td>658<br>MTD: 3,924</td><td>128</td></tr>
</table>
"""


def test_parse_litchi_row_reads_todays_figures_not_mtd():
    parsed = market_reference.parse_litchi_row(_TABLE_HTML)
    assert parsed == {"value_sold": 90200.0, "qty_sold": 312.0, "kg_sold": 658.0}


def test_parse_litchi_row_none_when_commodity_absent():
    """Most of the year Litchis isn't trading at all - some layouts drop
    the row entirely rather than show zeroes."""
    html = _TABLE_HTML.replace("LITCHIS", "LITCHIS-NOT-REALLY")
    assert market_reference.parse_litchi_row(html) is None


def test_sync_skips_a_zero_kg_day(client, monkeypatch):
    """Out of season the row (if present at all) sells 0 kg - never worth a
    row of its own, so the reference stays "no data" rather than a
    meaningless 0."""
    monkeypatch.setattr(market_reference, "_fetch_html",
                        lambda: _TABLE_HTML.replace("658", "0").replace("312", "0"))
    market_reference.sync_joburg_litchi_price()
    with Session(owner_engine) as owner:
        assert owner.exec(select(MarketDailyPrice)).all() == []


def test_sync_files_the_row_under_the_pages_own_trading_day(client, monkeypatch):
    """The page says which trading day it shows - over a weekend, still
    Friday's. Filed under that day, and a re-run rewrites the same row."""
    monkeypatch.setattr(market_reference, "_fetch_html", lambda: _TABLE_HTML)
    market_reference.sync_joburg_litchi_price()
    market_reference.sync_joburg_litchi_price()
    with Session(owner_engine) as owner:
        rows = owner.exec(select(MarketDailyPrice)).all()
    assert [(r.price_date, r.qty_sold, r.value_sold) for r in rows] == [(date(2025, 10, 24), 312.0, 90200.0)]


def test_sync_stores_nothing_without_a_readable_date(client, monkeypatch):
    """Under the wrong day it would be worse than missing."""
    monkeypatch.setattr(market_reference, "_fetch_html",
                        lambda: _TABLE_HTML.replace("24 October 2025", "sometime"))
    market_reference.sync_joburg_litchi_price()
    with Session(owner_engine) as owner:
        assert owner.exec(select(MarketDailyPrice)).all() == []
